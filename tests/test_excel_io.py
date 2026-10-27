"""Excel upload (template, multi-sheet workbooks, round trip) and Excel/PDF versions of every report."""
import io

import pytest
from django.contrib.auth.models import User
from django.test import Client
from openpyxl import Workbook, load_workbook

from ftth import importers
from ftth.exporters import excel, package
from ftth.exporters.common import stamp
from ftth.model import Design
from ftth.sample import sample_design


def workbook(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filled_template() -> bytes:
    """The downloadable template with a small network typed in."""
    wb = load_workbook(io.BytesIO(excel.survey_template_xlsx()))
    wb["Poles"].append(["P001", -1.29000, 36.82000, "wood", 9, "good", "existing"])
    wb["Poles"].append(["P002", -1.29040, 36.82000, "wood", 9, "leaning", "existing"])
    wb["Poles"].append(["P003", -1.29080, 36.82000, "concrete", 10, "good", "proposed"])
    wb["Equipment"].append(["Ex Closure 01", "ENCLOSURE", -1.29000, 36.82000, "P001", "existing"])
    wb["Equipment"].append(["FAT-01", "FAT", -1.29080, 36.82000, "P003", "proposed", 16])
    wb["Homes"].append(["H001", -1.29085, 36.82010, "Riverside Apartments", "MDU", 12])
    wb["Homes"].append(["H002", -1.29090, 36.81990, "Plot 4", "SDU", 1])
    wb["Splitters"].append(["", "FAT-01", 2, "1x8"])
    wb["Cables"].append(["CB01", "spur", 12, "Ex Closure 01", "FAT-01", "P001 > P002 > P003", "aerial", "proposed"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_template_has_a_sheet_per_record_kind():
    wb = load_workbook(io.BytesIO(excel.survey_template_xlsx()))
    assert wb.sheetnames == ["Instructions", "Poles", "Homes", "Equipment", "Splitters", "Cables"]
    assert [c.value for c in wb["Cables"][1]][:6] == ["ID", "Kind", "Cores", "From", "To", "Path"]


def test_filled_template_imports_every_sheet_in_one_go():
    d = Design()
    rep = importers.import_file(d, "survey.xlsx", filled_template())
    assert [p.id for p in d.poles] == ["P001", "P002", "P003"] and d.poles[1].condition == "leaning"
    assert {n.id: n.kind for n in d.nodes} == {"Ex Closure 01": "ENCLOSURE", "FAT-01": "FAT"}
    assert d.node("FAT-01").ports == 16 and d.node("FAT-01").pole == "P003"
    assert {h.id: (h.kind, h.units) for h in d.homes} == {"H001": ("mdu", 12), "H002": ("home", 1)}
    assert len(d.splitters) == 1 and d.splitters[0].stage == 2 and d.splitters[0].outputs == 8
    c = d.cable("CB01")
    assert (c.from_node, c.to_node, c.cores) == ("Ex Closure 01", "FAT-01", 12)
    assert [v.pole for v in c.path] == ["P001", "P002", "P003"]
    assert rep["added"] == 9 and any("Sheet 'Cables' as cables" in p for p in rep["problems"])


def test_sheets_without_names_are_detected_from_columns():
    d = Design()
    data = workbook({
        "Sheet1": [["Pole ID", "Latitude", "Longitude", "Height"], ["P9", -1.29, 36.82, 9]],
        "Sheet2": [["Building", "Latitude", "Longitude", "Viable HP"], ["Block A", -1.2901, 36.8201, 30]],
        "Sheet3": [["ID", "Latitude", "Longitude"], ["FAT-07", -1.2902, 36.8202]],
        "Notes": [["anything", "else"], ["free", "text"]],
    })
    rep = importers.import_file(d, "mixed.xlsx", data)
    assert [p.id for p in d.poles] == ["P9"] and d.homes[0].units == 30 and d.nodes[0].kind == "FAT"
    assert any("Sheet 'Notes' skipped" in p for p in rep["problems"])


def test_cable_with_unknown_pole_is_reported_not_guessed():
    d = Design()
    data = workbook({"Poles": [["ID", "Lat", "Lon"], ["P1", -1.29, 36.82]],
                     "Cables": [["ID", "Cores", "Path"], ["C1", 12, "P1 > P404"]]})
    rep = importers.import_file(d, "x.xlsx", data)
    assert d.cables == [] and any("P404" in p for p in rep["problems"])


def test_project_excel_export_round_trips():
    src = sample_design()
    data = excel.project_xlsx(src, stamp(src, 1))
    d = Design()
    importers.import_file(d, "project.xlsx", data, status="proposed")
    assert {p.id for p in d.poles} == {p.id for p in src.poles}
    assert {n.id: n.kind for n in d.nodes} == {n.id: n.kind for n in src.nodes}
    assert sum(h.units for h in d.homes) == sum(h.units for h in src.homes)
    assert {c.id for c in d.cables} == {c.id for c in src.cables}
    assert len(d.splitters) == len(src.splitters)


def test_neighbouring_buildings_in_one_sheet_stay_separate():
    d = Design()
    data = workbook({"Homes": [["ID", "Latitude", "Longitude", "Viable HP"],
                               ["Block A", -1.29000, 36.82000, 10], ["Block B", -1.29003, 36.82003, 8]]})
    importers.import_file(d, "h.xlsx", data)
    assert [h.id for h in d.homes] == ["Block A", "Block B"]


def test_cable_bends_that_are_not_poles_come_from_coordinates():
    d = Design()
    data = workbook({"Poles": [["ID", "Lat", "Lon"], ["P1", -1.29, 36.82], ["P2", -1.291, 36.821]],
                     "Cables": [["ID", "Cores", "Path", "Path coordinates"],
                                ["C1", 12, "P1 > pt > P2", "-1.29,36.82; -1.2905,36.8203; -1.291,36.821"]]})
    importers.import_file(d, "c.xlsx", data)
    assert [(v.pole, round(v.lat, 4)) for v in d.cable("C1").path] == [("P1", -1.29), (None, -1.2905), ("P2", -1.291)]


def test_explicit_target_still_reads_first_sheet():
    d = Design()
    data = workbook({"Data": [["ID", "Latitude", "Longitude"], ["X1", -1.29, 36.82]]})
    importers.import_file(d, "x.xlsx", data, target="poles")
    assert [p.id for p in d.poles] == ["X1"]


@pytest.mark.parametrize("deliv", ["pole_summary", "bom", "boq", "splicing", "report", "home_list", "qa"])
def test_every_report_prints_as_pdf_and_excel(deliv):
    d = sample_design()
    fmts = package.DELIVERABLES[deliv]
    assert {"pdf", "xlsx"} <= set(fmts)
    st = stamp(d, 1)
    assert package.build(d, deliv, "pdf", st, None)[:5] == b"%PDF-"
    wb = load_workbook(io.BytesIO(package.build(d, deliv, "xlsx", st, None)))
    assert wb.sheetnames


@pytest.mark.django_db
def test_template_download_and_web_import(tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    u = User.objects.create_user("ann", password="pw-12345-long")
    c = Client()
    c.force_login(u)
    r = c.get("/api/survey-template.xlsx")
    assert r.status_code == 200 and r["Content-Type"].startswith("application/vnd.openxmlformats")
    pid = c.post("/api/projects", {"name": "X"}, content_type="application/json").json()["id"]
    f = io.BytesIO(filled_template())
    f.name = "survey.xlsx"
    st = c.post(f"/api/projects/{pid}/import", {"file": f, "target": "auto", "status": "existing"}).json()
    assert st["import"]["added"] == 9 and len(st["design"]["cables"]) == 1
    assert c.get(f"/api/projects/{pid}/export/qa.xlsx").status_code == 200
    assert c.get(f"/api/projects/{pid}/export/report.xlsx").status_code == 200
