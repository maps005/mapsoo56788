"""Sectors: a company plans fibre, water supply, pipelines, electricity, street lighting or sanitation networks."""
import io
import zipfile

from openpyxl import load_workbook

import pytest
from django.contrib.auth.models import User
from django.test import Client
from pypdf import PdfReader

from ftth import calc, sectors, validate
from ftth.exporters import excel, package, pdf, report
from ftth.exporters.common import stamp
from ftth.model import Cable, Vertex
from ftth.sample import sample_design
from planner.models import Company, Membership, Profile


@pytest.fixture(autouse=True)
def media(tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path


def water_company_user():
    u = User.objects.create_user("aqua", password="pw-12345-long")
    co = Company.objects.create(name="Aqua Works", sector="water")
    Membership.objects.create(user=u, company=co, role="owner")
    Profile.objects.update_or_create(user=u, defaults={"active_company": co})
    c = Client()
    c.force_login(u)
    return u, co, c


@pytest.mark.parametrize("key", ["water", "pipelines", "electricity", "lighting", "sanitation"])
def test_each_sector_names_its_materials_and_report(key):
    d = sample_design()
    for f in ("introduction", "findings", "proposed_solution", "special_requirements", "recommendations"):
        setattr(d.info, f, "")  # the demo's own text is about fibre
    sectors.apply(d, key)
    p = sectors.SECTORS[key]
    assert d.info.sector == key and d.rules.cable_code == p["rules"]["cable_code"]
    assert {m.code for m in d.materials} == {m.code for m in p["materials"]}
    secs = report.sections(d)
    text = " ".join(x for s in secs for x in s.paragraphs) + " ".join(s.title for s in secs)
    assert "FTTH" not in text and "homepass" not in text and "splitter" not in text.lower()
    out = " ".join(pg.extract_text() for pg in PdfReader(io.BytesIO(pdf.report_pdf(d, stamp(d, 1), None))).pages)
    assert p["report_title"].upper() in out and p["kinds"]["FAT"] in out


def test_water_pipe_sizes_price_from_the_catalogue():
    d = sample_design()
    sectors.apply(d, "water")
    d.cables = [Cable(id="M1", kind="main", cores=160, status="proposed", install="underground",
                      path=[Vertex(lat=-1.29, lon=36.82), Vertex(lat=-1.29, lon=36.8209)])]
    line = next(ln for ln in calc.bom(d) if ln.code == "PIPE-HDPE-DN160")
    assert line.description == "HDPE pipe PN10, DN160 mm" and line.qty > 90


@pytest.mark.django_db
def test_new_projects_follow_the_company_sector_and_can_switch():
    u, co, c = water_company_user()
    st = c.post("/api/projects", {"name": "Kiambu water"}, content_type="application/json").json()
    assert st["sector"]["key"] == "water" and st["sector"]["kinds"]["FAT"] == "Meter / service connection"
    assert st["sector"]["fibre_tools"] is False and st["company"]["sector"] == "water"
    assert any(m["code"] == "VALVE-CHAMBER" for m in st["design"]["materials"])
    r = c.post(f"/api/projects/{st['id']}/apply-sector", {"sector": "electricity"}, content_type="application/json")
    assert r.status_code == 200 and r.json()["sector"]["key"] == "electricity"
    assert c.post(f"/api/projects/{st['id']}/apply-sector", {"sector": "gold"},
                  content_type="application/json").status_code == 400
    assert [x["key"] for x in c.get("/api/sectors").json()] == list(sectors.SECTORS)


@pytest.mark.django_db
def test_company_page_and_admin_set_the_sector():
    owner = User.objects.create_superuser("boss", password="pw-12345-long")
    co = Company.objects.create(name="Grid Co")
    Membership.objects.create(user=owner, company=co, role="owner")
    c = Client()
    c.force_login(owner)
    c.post("/company/", {"c": co.id, "action": "details", "name": "Grid Co", "currency": "KES",
                         "sector": "electricity", "design_standard": "system"})
    assert Company.objects.get(pk=co.id).sector == "electricity"
    page = c.get("/admin/planner/company/").content
    assert b"Electricity distribution" in page and b"sector__exact" in page


FIBRE_HARDWARE = ("J-HOOK", "UPB", "TENSION-CLAMP", "FOC-", "SPL-1X", "ENC-SPLICE", "FAT-16P", "STEEL-STRAP")
BURIED = [k for k, v in sectors.SECTORS.items() if v["install"] == "underground"]
# network sectors other than fibre (construction designs buildings: tests/test_construction.py)
NETWORKS = [k for k, v in sectors.SECTORS.items() if k != "fibre" and not v.get("construction")]


def sector_design(key):
    d = sample_design()
    for f in ("introduction", "findings", "proposed_solution", "special_requirements", "recommendations"):
        setattr(d.info, f, "")
    sectors.apply(d, key)
    return d


@pytest.mark.parametrize("key", NETWORKS)
def test_sector_boq_has_only_the_sectors_items_and_sizes(key):
    d = sector_design(key)
    p = sectors.SECTORS[key]
    assert all(c.cores in p["sizes"] for c in d.cables)  # fibre 12C / 48C became the sector's sizes
    lines = calc.bom(d)
    codes = {ln.code for ln in lines}
    assert not [c for c in codes if c.startswith(FIBRE_HARDWARE)]
    assert all(ln.rate is not None for ln in lines), [ln.code for ln in lines if ln.rate is None]  # all in catalogue
    assert all(cl.slack_m == 0 and cl.slack_loops == 0 for cl in calc.cable_lengths(d))
    if key in BURIED:
        assert all(cl.sag_m == 0 for cl in calc.cable_lengths(d)) and "CIVIL-TRENCH" in codes
        assert all(c.install == "underground" for c in d.cables if c.status == "proposed")
        rows = calc.pole_summary(d)
        assert sum(r.jhooks + r.tension + r.upb + r.ext_arms for r in rows) == 0
        assert any(r.cables for r in rows)  # buried lines are still listed at their manholes / marker posts
    msgs = " ".join(i["message"] for i in validate.validate(d))
    assert " FAT" not in msgs and "Home H" not in msgs


@pytest.mark.parametrize("key", NETWORKS)
def test_every_deliverable_builds_in_the_sectors_words(key, tmp_path):
    d = sector_design(key)
    p = sectors.SECTORS[key]
    st = stamp(d, 1)
    with zipfile.ZipFile(io.BytesIO(package.build_package(d, st, tmp_path))) as z:
        assert z.namelist()  # KML/KMZ, DXF, LLD, schedules, BOQ, report: every file builds
    ws = load_workbook(io.BytesIO(excel.pole_summary_xlsx(d, st))).active
    heads = [c.value for row in ws.iter_rows(min_row=1, max_row=6) for c in row if isinstance(c.value, str)]
    assert f"{p['support']} ID" in heads and "J-Hooks" not in heads and p["kinds"]["ENCLOSURE"] in heads
    text = " ".join(pg.extract_text() for pg in PdfReader(io.BytesIO(pdf.pole_summary_pdf(d, st))).pages)
    assert f"{p['support']} Schedule" in text and "J-Hooks" not in text


def test_electricity_uses_overhead_hardware_stays_and_transformer_ratings():
    d = sector_design("electricity")
    codes = {ln.code: ln.qty for ln in calc.bom(d)}
    assert codes["DEAD-END-CLAMP"] and codes["STAY-SET"] == -(-codes["DEAD-END-CLAMP"] // 2)
    assert "TX-100KVA" in codes  # the PON node became a 100 kVA transformer
    for h in d.homes:
        h.units = 40
    msgs = [i for i in validate.validate(d) if i["code"] == "tx_overload"]
    assert msgs and "exceeds the transformer capacity 100 kVA" in msgs[0]["message"]
    for h in d.homes:
        h.details["load_kva"] = 0.2  # a recorded load replaces the default after-diversity demand
    assert not [i for i in validate.validate(d) if i["code"] == "tx_overload"]


def test_lighting_columns_get_luminaires_from_their_details_and_spacing_is_checked():
    d = sector_design("lighting")
    new = [p for p in d.poles if p.status == "proposed"]
    new[0].details.update({"lamps": 2, "lamp_w": "120"})
    q = {ln.code: ln.qty for ln in calc.bom(d)}
    assert q["LIGHT-COLUMN"] == len(new) and q["LUMINAIRE-120W"] == 2
    assert q.get("LUMINAIRE-60W", 0) >= len(new) - 1  # other new columns: one 60 W luminaire by default
    assert any(i["code"] == "column_spacing" for i in validate.validate(d))  # demo poles are ~120 m apart


def test_sewer_falls_and_manhole_levels_are_checked():
    d = sector_design("sanitation")
    c = d.cables[0]
    feats = [f for f in ((d.node(v.node) if v.node else None) or (d.pole(v.pole) if v.pole else None)
                         for v in c.path) if f]
    for k, f in enumerate(feats):
        f.details["invert_level_m"] = 100 - k * 2  # good fall all the way
    assert not [i for i in validate.validate(d) if i["code"].startswith("sewer")]
    feats[1].details["invert_level_m"] = 100 - 0.01  # almost flat between the first two manholes
    assert any(i["code"] == "sewer_gradient" for i in validate.validate(d))
    feats[1].details["invert_level_m"] = 101  # higher than the manhole upstream
    assert any(i["code"] == "sewer_adverse" for i in validate.validate(d))
    feats[0].details["cover_level_m"] = 99
    assert any(i["code"] == "manhole_levels" and i["ref"] == feats[0].id for i in validate.validate(d))


def test_qa_rewording_keeps_feature_ids():
    words = sectors.wording("water")
    assert sectors.reword("Home H001 is not served by a FAT (FAT-01, POLE-NEW)", words) == \
        "Property H001 is not served by a service connection (FAT-01, POLE-NEW)"
    assert sectors.reword("a FAT", sectors.wording("pipelines")) == "an offtake"


def test_switching_back_to_fibre_keeps_splitters_and_connections():
    d = sample_design()
    n_split, n_conn = len(d.splitters), len(d.connections)
    sectors.apply(d, "water")
    assert not [ln for ln in calc.bom(d) if ln.code.startswith("SPL-")]  # ignored, not deleted
    sectors.apply(d, "fibre")
    assert (len(d.splitters), len(d.connections)) == (n_split, n_conn)


@pytest.mark.django_db
def test_lighting_company_projects_record_details_and_keep_their_sector():
    u = User.objects.create_user("lux", password="pw-12345-long")
    co = Company.objects.create(name="Lux Lighting", sector="lighting")
    Membership.objects.create(user=u, company=co, role="owner")
    Profile.objects.update_or_create(user=u, defaults={"active_company": co})
    c = Client()
    c.force_login(u)
    st = c.post("/api/projects", {}, content_type="application/json").json()
    assert st["design"]["info"]["name"] == "New street lighting project"
    assert st["sector"]["support"] == "Lighting column" and st["sector"]["install"] == "underground"
    assert any(f["key"] == "lamp_w" for f in st["sector"]["details"]["pole"])
    d = st["design"]
    d["poles"].append({"id": "LC1", "lat": -1.29, "lon": 36.82, "status": "proposed",
                       "details": {"lamps": 2, "lamp_w": "90", "bracket": "Double arm"}})
    st = c.put(f"/api/projects/{st['id']}", d, content_type="application/json", HTTP_X_BASE_HASH=st["hash"]).json()
    assert st["design"]["poles"][-1]["details"]["lamp_w"] == "90"
    calc_ = c.get(f"/api/projects/{st['id']}/calc").json()
    assert {"code": "LUMINAIRE-90W"}.items() <= next(x for x in calc_["bom"] if x["code"] == "LUMINAIRE-90W").items()
    # the company changing sector later does not rewrite its existing projects
    co.sector = "water"
    co.save()
    assert c.get(f"/api/projects/{st['id']}").json()["sector"]["key"] == "lighting"
