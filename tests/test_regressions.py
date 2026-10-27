"""Regression tests for bugs found by the 2026-10-26 agent test round (UI, API/security, engineering logic)."""
import io
import zipfile

import pytest
from django.contrib.auth.models import Permission, User
from django.test import Client

from ftth import calc, importers, splicing, validate
from ftth.exporters import package
from ftth.exporters.common import photo_path, stamp
from ftth.exporters.kml import build_kml
from ftth.model import Cable, Design, Node, Pole, Vertex
from ftth.sample import sample_design


def kml(*folders):
    return ('<?xml version="1.0"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document>'
            + "".join(folders) + "</Document></kml>").encode()


def folder(name, *pms):
    return f"<Folder><name>{name}</name>{''.join(pms)}</Folder>"


def pt(name, lat, lon):
    return f"<Placemark><name>{name}</name><Point><coordinates>{lon},{lat},0</coordinates></Point></Placemark>"


def line(name, pts):
    c = " ".join(f"{lon}, {lat}, 0" for lat, lon in pts)  # spaces after commas, as some tools write
    return f"<Placemark><name>{name}</name><LineString><coordinates>{c}</coordinates></LineString></Placemark>"


# ---- importer ------------------------------------------------------------------
def test_identical_pole_coordinates_do_not_crash_import():
    data = kml(folder("POLES", pt("A", -1.3, 36.8), pt("B", -1.3, 36.8)),
               folder("CABLES", line("12C", [(-1.3, 36.8), (-1.3, 36.801)])))
    d = Design()
    importers.import_file(d, "a.kml", data, "auto")
    assert len(d.cables) == 1 and len(d.cables[0].path) == 2
    rep = importers.import_file(d, "a.kml", data, "auto")  # importing twice warns
    assert any("imported again" in p for p in rep["problems"])


def test_pon_folder_does_not_turn_fats_into_pons():
    d = Design()
    data = kml(folder("PON 1", folder("FATs", pt("FAT-01", -1.3, 36.8))))
    importers.import_file(d, "a.kml", data, "auto")
    assert d.nodes[0].kind == "FAT"


def test_kmz_zip_bomb_is_refused():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("doc.kml", b"<kml>" + b" " * (importers.MAX_KML_BYTES + 10) + b"</kml>")
    with pytest.raises(ValueError):
        importers.parse_kml(buf.getvalue())


def test_csv_uses_status_ports_kind_and_comma_decimals():
    csv = (b"Survey sheet\nID,Coordinates,Kind,Status,Ports\n"
           b'PR FAT-03,"-1,3 36,8",FAT,proposed,16\nX1,-1.3;36.8,slack,existing,\nX1,-1.31 36.81,,,\n')
    d = Design()
    rep = importers.import_file(d, "n.csv", csv, "nodes")
    n = d.nodes[0]
    assert (n.kind, n.status, n.ports, n.lat, n.lon) == ("FAT", "proposed", 16, -1.3, 36.8)
    assert d.nodes[1].kind == "SLACK" and len(d.nodes) == 3  # duplicate ID kept and reported
    assert any("X1 also appears" in p for p in rep["problems"])


def test_kml_round_trip_preserves_quantities():
    d = sample_design()
    calc.assign_homes(d)
    splicing.auto_allocate(d)
    d.cables[2].path[0].crossing = True
    d2 = Design(materials=d.materials)
    importers.import_file(d2, "x.kml", build_kml(d, stamp(d, 1)).encode(), "auto")
    assert {l.code: l.qty for l in calc.bom(d)} == {l.code: l.qty for l in calc.bom(d2)}
    assert calc.summary(d)["homepasses"] == calc.summary(d2)["homepasses"]
    assert sorted(c.id for c in d2.cables) == sorted(c.id for c in d.cables)


# ---- calculations / allocation ----------------------------------------------------
def _straight_cable(d, pts, cid="CX"):
    for i, (la, lo) in enumerate(pts):
        if not d.pole(f"Q{i}"):
            d.poles.append(Pole(id=f"Q{i}", lat=la, lon=lo, status="existing"))
    return Cable(id=cid, kind="spur", cores=12, path=[Vertex(lat=la, lon=lo, pole=f"Q{i}") for i, (la, lo) in enumerate(pts)])


def test_duplicate_route_points_do_not_create_corners():
    d = Design()
    pts = [(-1.3, 36.8), (-1.3, 36.8004), (-1.3, 36.8008), (-1.3, 36.8012)]
    c = _straight_cable(d, pts)
    c.path.insert(2, c.path[1].model_copy())  # the same pole drawn twice
    d.cables.append(c)
    rows = {r.pole: r for r in calc.pole_summary(d)}
    assert rows["Q1"].tension == 0 and rows["Q1"].jhooks == 1


def test_one_point_cable_gets_no_accessories():
    d = Design(poles=[Pole(id="P1", lat=-1.3, lon=36.8)])
    d.cables.append(Cable(id="C1", path=[Vertex(lat=-1.3, lon=36.8, pole="P1")]))
    assert all(l.qty == 0 for l in calc.bom(d) if l.code in ("TENSION-CLAMP", "UPB"))


def test_fat_added_after_allocation_is_fed():
    d = sample_design()
    calc.assign_homes(d)
    splicing.auto_allocate(d)
    jb = d.node("JB01")
    d.nodes.append(Node(id="FAT03", kind="FAT", lat=jb.lat + 0.0005, lon=jb.lon, ports=16))
    d.cables.append(Cable(id="C099", kind="spur", cores=12, from_node="JB01", to_node="FAT03",
                          path=[Vertex(lat=jb.lat, lon=jb.lon, node="JB01", pole=jb.pole),
                                Vertex(lat=jb.lat + 0.0005, lon=jb.lon, node="FAT03")]))
    res = splicing.auto_allocate(d)
    assert res["issues"] == []
    s = next(s for s in d.splitters if s.node == "FAT03")
    assert splicing.trace_upstream(d, s.id)["ok"]


def test_unconnected_pop_is_not_chosen_as_root():
    d = sample_design()
    d.nodes.append(Node(id="POP00", kind="POP", lat=-1.40, lon=36.90, status="existing"))
    res = splicing.auto_allocate(d)
    assert res["issues"] == []


def test_unfed_fat_blocks_approval_once_allocated():
    d = sample_design()
    calc.assign_homes(d)
    splicing.auto_allocate(d)
    d.connections = [c for c in d.connections if c.b.ref != "S2-FAT02" and c.a.ref != "S2-FAT02"]
    errs = [i for i in validate.validate(d) if i["severity"] == "error"]
    assert any(i["code"] == "fat_unfed" for i in errs)


# ---- exporters ------------------------------------------------------------------
def test_markup_and_control_characters_in_user_text():
    d = sample_design()
    d.info.area = '<img src="http://127.0.0.1:9/should-not-be-fetched"/> R&D'
    d.poles[0].notes = "Check <span> clearance \x0b =1+1"
    st = stamp(d, 1)
    for k, fmts in package.DELIVERABLES.items():
        for f in fmts:
            package.build(d, k, f, st, None)  # must not raise or fetch anything


def test_excel_never_writes_user_text_as_formula():
    from openpyxl import load_workbook
    d = sample_design()
    d.poles[0].notes = "=HYPERLINK(\"http://evil\")"
    wb = load_workbook(io.BytesIO(package.build(d, "project", "xlsx", stamp(d, 1), None)))
    cell = wb["Poles"]["H5"]
    assert cell.data_type == "s"


def test_photo_paths_cannot_leave_the_project_folder(tmp_path):
    (tmp_path / "secret.png").write_bytes(b"x")
    assert photo_path(tmp_path / "proj", "../secret.png") is None
    assert photo_path(tmp_path / "proj", str(tmp_path / "secret.png")) is None
    with pytest.raises(Exception):
        Design.model_validate({"photos": [{"id": "p", "file": "../a/b.png"}]})


# ---- web layer --------------------------------------------------------------------
@pytest.fixture(autouse=True)
def media(tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path


def login(username, perms=()):
    u = User.objects.create_user(username, password="pw-12345-long")
    for p in perms:
        u.user_permissions.add(Permission.objects.get(codename=p))
    c = Client()
    c.force_login(u)
    return c


@pytest.mark.django_db
def test_put_requires_base_hash_and_bad_bodies_are_400():
    c = login("ann")
    st = c.post("/api/projects", {"sample": True}, content_type="application/json").json()
    pid = st["id"]
    assert c.put(f"/api/projects/{pid}", st["design"], content_type="application/json").status_code == 428
    for url, body in ((f"/api/projects", [1, 2]), (f"/api/projects/{pid}/share", {"username": 7}),
                      (f"/api/projects/{pid}/revisions", {"note": 5})):
        assert c.post(url, body, content_type="application/json").status_code == 400
    assert c.post(f"/api/projects/{pid}/assign-homes?max_drop_m=abc").status_code == 400


@pytest.mark.django_db
def test_survey_approval_is_checked_and_reset_by_survey_edits():
    c = login("ann")
    st = c.post("/api/projects", {"sample": True}, content_type="application/json").json()
    pid, d = st["id"], st["design"]
    d["info"]["survey_status"] = "approved"
    st = c.put(f"/api/projects/{pid}", d, content_type="application/json", HTTP_X_BASE_HASH=st["hash"]).json()
    assert st["design"]["info"]["survey_approved_by"] == "ann"
    d = st["design"]
    d["poles"][0]["lat"] += 0.001
    st = c.put(f"/api/projects/{pid}", d, content_type="application/json", HTTP_X_BASE_HASH=st["hash"]).json()
    assert st["design"]["info"]["survey_status"] == "draft"
    d = st["design"]
    d["poles"][0]["lat"] = None  # survey error
    d["info"]["survey_status"] = "approved"
    r = c.put(f"/api/projects/{pid}", d, content_type="application/json", HTTP_X_BASE_HASH=st["hash"])
    assert r.status_code == 409


@pytest.mark.django_db
def test_restore_is_recorded_and_returns_to_draft():
    c = login("eve", perms=["approve_design"])
    st = c.post("/api/projects", {"sample": True}, content_type="application/json").json()
    pid = st["id"]
    c.post(f"/api/projects/{pid}/assign-homes")
    c.post(f"/api/projects/{pid}/allocate")
    c.post(f"/api/projects/{pid}/revisions", {"note": "issue", "design_status": "approved"}, content_type="application/json")
    st = c.post(f"/api/projects/{pid}/revisions/1/restore").json()
    assert st["design"]["info"]["design_status"] == "draft"
    revs = c.get(f"/api/projects/{pid}/revisions").json()
    assert revs[0]["note"] == "Restored revision 1"


@pytest.mark.django_db
def test_only_affected_outputs_need_regenerating():
    c = login("ann")
    st = c.post("/api/projects", {"sample": True}, content_type="application/json").json()
    pid = st["id"]
    for k in ("splicing", "pole_summary"):
        c.get(f"/api/projects/{pid}/export/{k}.pdf")
    st = c.get(f"/api/projects/{pid}").json()
    d = st["design"]
    d["poles"][0]["height_m"] = 11
    st = c.put(f"/api/projects/{pid}", d, content_type="application/json", HTTP_X_BASE_HASH=st["hash"]).json()
    assert st["outputs"]["pole_summary"]["stale"] is True
    assert st["outputs"]["splicing"]["stale"] is False


@pytest.mark.django_db
def test_login_is_throttled_and_pages_send_csp():
    User.objects.create_user("bob", password="right-password-123")
    c = Client()
    for _ in range(10):
        c.post("/accounts/login/", {"username": "bob", "password": "wrong"})
    r = c.post("/accounts/login/", {"username": "bob", "password": "right-password-123"})
    assert r.status_code == 200 and b"Too many failed" in r.content
    assert "script-src 'self'" in c.get("/accounts/login/")["Content-Security-Policy"]


# ---- approved-project features (2026-10-26) ----------------------------------------
def test_home_list_merges_with_kmz_homes_and_skips_section_rows():
    from openpyxl import Workbook
    d = Design()
    importers.import_file(d, "a.kml", kml(folder("HOMES", folder("MDU", pt("MDU 121 UC", -1.3, 36.8)))), "auto")
    assert d.homes[0].units == 0  # no count in the name: not counted
    wb = Workbook()
    ws = wb.active
    ws.append(["REGION", "ZONE", "CLUSTER", "GREEN/BROWN FIELD", "NAME", "NATURE OF BUILDING", "VIABLE HP",
               "COMMERCIAL", "LONGITUDE", "LATITUDE", "ACQUIRED"])
    ws.append(["MULTI DWELLING UNITS (MDUs)"])
    ws.append(["NAIROBI", "KASARANI", "INFINITY", "GREEN FIELD", "MDU 121 UC", "MDU", 18, None, 36.8, -1.3, "ACQUIRED"])
    ws.append([None, None, None, None, None, "TOTAL", 18])
    buf = io.BytesIO()
    wb.save(buf)
    importers.import_file(d, "home list.xlsx", buf.getvalue(), "homes")
    assert len(d.homes) == 1 and d.homes[0].units == 18 and d.homes[0].acquired
    assert d.info.region == "Nairobi" and d.info.cluster == "Infinity"
    out = package.build(d, "home_list", "xlsx", stamp(d, 1), None)
    d2 = Design()
    importers.import_file(d2, "hl.xlsx", out, "homes")  # the exported list imports back
    assert sum(h.units for h in d2.homes) == 18


def test_combined_splitter_label_existing_cable_folder_and_trunking():
    data = kml(folder("SPLICE ENCLOSURE", pt("PR SC-01", -1.3, 36.8)),
               folder("FATs", pt("PR IND FAT-01", -1.3, 36.801)),
               folder("SPLITTERS", folder("SECONDARY SPLITTERS", pt("1X8 SPL 1A &amp; 2A", -1.3, 36.801))),
               folder("CABLES", folder("EXISTING", line("24C 3", [(-1.31, 36.8), (-1.3, 36.8)])),
                      folder("12C ADSS FOC", line("12C", [(-1.3, 36.8), (-1.3, 36.801)]))),
               folder("PLASTIC TRUNKING", line("Trunking", [(-1.3, 36.801), (-1.3001, 36.801)])))
    d = Design()
    importers.import_file(d, "a.kml", data, "auto")
    assert sorted(s.id for s in d.splitters) == ["1X8 SPL 1A", "1X8 SPL 2A"]
    status = {c.cores: c.status for c in d.cables}
    assert status == {24: "existing", 12: "proposed"}
    assert len(d.ducts) == 1 and any(l.code == "TRUNKING" for l in calc.bom(d))
    d2 = Design()
    importers.import_file(d2, "x.kml", build_kml(d, stamp(d, 1)).encode(), "auto")
    assert len(d2.ducts) == 1  # trunking survives the KML round trip


def test_large_mdu_shares_nearby_fats():
    d = Design(homes=[], nodes=[])
    from ftth.model import Home
    d.nodes += [Node(id="F1", kind="FAT", lat=-1.3, lon=36.8, ports=24),
                Node(id="F2", kind="FAT", lat=-1.3, lon=36.8003, ports=24)]
    d.homes.append(Home(id="MDU", kind="mdu", units=36, lat=-1.3, lon=36.8001))
    assert calc.assign_homes(d)["unassigned"] == []
    assert d.homes[0].fats == ["F1", "F2"] and calc.fat_load(d) == {"F1": 24, "F2": 12}
    assert not [i for i in validate.validate(d) if i["code"] == "fat_capacity"]


@pytest.mark.django_db
def test_kmz_photos_are_imported_into_the_project(tmp_path):
    from PIL import Image
    img = io.BytesIO()
    Image.new("RGB", (6, 6), "blue").save(img, "JPEG")
    doc = kml(folder("POLES", '<Placemark><name>P1</name><description><![CDATA[<img src="files/p1.jpg">]]>'
                              '</description><Point><coordinates>36.8,-1.3,0</coordinates></Point></Placemark>'))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("doc.kml", doc)
        z.writestr("files/p1.jpg", img.getvalue())
    c = login("ann")
    pid = c.post("/api/projects", {"name": "A"}, content_type="application/json").json()["id"]
    up = io.BytesIO(buf.getvalue())
    up.name = "survey.kmz"
    st = c.post(f"/api/projects/{pid}/import", {"file": up, "target": "auto"}).json()
    photos = st["design"]["photos"]
    assert len(photos) == 1 and photos[0]["feature"] == "P1"
    assert c.get(f"/api/projects/{pid}/photos/{photos[0]['file']}").status_code == 200


def test_survey_facts_are_notes_not_warnings():
    """Existing GIS cables, blank homepasses and a backbone stub left for a future phase are not design problems."""
    d = sample_design()
    p = d.poles[0]
    d.cables.append(Cable(id="EX-48C", kind="main", cores=48, status="existing",
                          path=[Vertex(lat=p.lat, lon=p.lon), Vertex(lat=p.lat + 0.001, lon=p.lon)]))
    d.homes[0].units = 0
    stub = Cable(id="STUB", kind="main", cores=48, status="proposed", from_node=d.nodes[0].id,
                 path=[Vertex(lat=d.nodes[0].lat, lon=d.nodes[0].lon, node=d.nodes[0].id),
                       Vertex(lat=p.lat, lon=p.lon, pole=p.id)])
    d.cables.append(stub)
    iss = validate.validate(d)
    mine = [i for i in iss if i["ref"] in ("EX-48C", "STUB", d.homes[0].id)]
    assert mine and all(i["severity"] == "info" for i in mine)
    assert validate.counts(iss)["notes"] >= 2
    stub.kind = "spur"  # an open spur is still a real warning
    assert any(i["code"] == "cable_unterminated" and i["ref"] == "STUB" for i in validate.validate(d))
