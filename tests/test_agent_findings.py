"""Fixes for the problems found by the agent planner test on 2026-10-27 (agent-test/agent-report.md)."""
import io

import pytest
from django.contrib.auth.models import Permission, User
from django.test import Client
from openpyxl import load_workbook

from ftth import autodesign, calc, importers
from ftth.exporters import excel, report
from ftth.geo import haversine
from ftth.model import Cable, Design, Home, Node, Pole, Vertex
from ftth.sample import sample_design


def client(perms=()):
    u = User.objects.create_user("pat", password="pw-12345-long")
    for p in perms:
        u.user_permissions.add(Permission.objects.get(codename=p))
    c = Client()
    c.force_login(u)
    return c


def put(c, pid, design, base):
    return c.put(f"/api/projects/{pid}", design, content_type="application/json", HTTP_X_BASE_HASH=base)


@pytest.mark.django_db
def test_design_work_keeps_the_survey_approved_but_survey_edits_do_not(tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    c = client()
    st = c.post("/api/projects", {"sample": True}, content_type="application/json").json()
    pid, d = st["id"], st["design"]
    d["info"]["survey_status"] = "approved"
    st = put(c, pid, d, st["hash"]).json()
    assert st["design"]["info"]["survey_status"] == "approved"
    st = c.post(f"/api/projects/{pid}/auto-design", {}, content_type="application/json").json()
    d = st["design"]
    d["nodes"][-1]["notes"] = "moved by the planner"          # design edits ...
    d["homes"][0]["fat"] = None
    st = put(c, pid, d, st["hash"]).json()
    assert st["design"]["info"]["survey_status"] == "approved"  # ... keep the survey approval
    d = st["design"]
    d["homes"][0]["units"] += 1                                # a survey edit returns it to draft
    st = put(c, pid, d, st["hash"]).json()
    assert st["design"]["info"]["survey_status"] == "draft"


@pytest.mark.django_db
def test_design_with_unpriced_boq_items_cannot_be_approved(tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    c = client(perms=["approve_design"])
    st = c.post("/api/projects", {"sample": True}, content_type="application/json").json()
    pid = st["id"]
    c.post(f"/api/projects/{pid}/assign-homes")
    st = c.post(f"/api/projects/{pid}/allocate").json()
    d = st["design"]
    d["materials"] = [m for m in d["materials"] if m["code"] != "ROAD-X"]
    put(c, pid, d, st["hash"])
    r = c.post(f"/api/projects/{pid}/revisions", {"note": "issue", "design_status": "approved"},
               content_type="application/json")
    assert r.status_code == 409 and "ROAD-X" in r.json()["detail"]


def test_template_name_column_is_the_building_name():
    wb = load_workbook(io.BytesIO(excel.survey_template_xlsx()))
    wb["Homes"].append(["H001", -1.29, 36.82, "Riverside Apartments", "MDU", 24])
    buf = io.BytesIO()
    wb.save(buf)
    d = Design()
    importers.import_file(d, "s.xlsx", buf.getvalue())
    assert d.homes[0].id == "H001" and d.homes[0].name == "Riverside Apartments"


def street(n=8, step=40.0, lat0=-1.29, lon0=36.82):
    """Poles along a straight street, an existing closure at the start."""
    d = Design()
    for k in range(n):
        d.poles.append(Pole(id=f"P{k:02d}", lat=lat0, lon=lon0 + k * step / 111_320, status="existing"))
    d.nodes.append(Node(id="EX SC", kind="ENCLOSURE", lat=lat0, lon=lon0, pole="P00", status="existing"))
    return d


def test_zero_homepass_premises_get_no_fat():
    d = street()
    d.homes += [Home(id="MDU A", lat=-1.2903, lon=36.8215, units=20), Home(id="Shop", lat=-1.2902, lon=36.8214, units=0)]
    autodesign.auto_design(d)
    h = {x.id: x for x in d.homes}
    assert h["MDU A"].fat and not h["Shop"].fat
    assert calc.assign_homes(d)["unassigned"] == []


def test_auto_design_goes_pole_to_pole_without_skipping():
    d = street(n=8, step=30.0)  # 30 m spacing: P00-P02 (60 m) is within the max span but passes P01
    d.homes.append(Home(id="Block", lat=-1.2903, lon=36.82 + 7 * 30 / 111_320, units=24))
    autodesign.auto_design(d)
    main = [c for c in d.cables if c.kind == "main"]
    poles = {p.id: p for p in d.poles}
    for c in main + [c for c in d.cables if c.kind == "spur"]:
        ps = [v.pole for v in c.path if v.pole]
        for a, b in zip(ps, ps[1:]):
            if a != b:
                assert haversine(poles[a].lat, poles[a].lon, poles[b].lat, poles[b].lon) < 45, (c.id, a, b)


def test_auto_design_follows_the_existing_pole_line_and_survey_notes():
    d = street(n=6, step=50.0)
    # a parallel, slightly shorter line of poles 30 m north; the existing cable runs on the southern line
    for k in range(1, 5):
        d.poles.append(Pole(id=f"N{k}", lat=-1.29 + 30 / 110_540, lon=36.82 + k * 49 / 111_320, status="existing"))
    d.cables.append(Cable(id="EX 24C", kind="spur", cores=24, status="existing",
                          path=[Vertex(lat=p.lat, lon=p.lon, pole=p.id) for p in d.poles[:6]]))
    d.homes.append(Home(id="Block", lat=-1.2902, lon=36.82 + 250 / 111_320, units=24))
    autodesign.auto_design(d)
    used = {v.pole for c in d.cables if c.status == "proposed" for v in c.path}
    assert {"P01", "P02", "P03", "P04"} <= used and not used & {"N1", "N2", "N3", "N4"}


def test_survey_note_asks_for_more_fats():
    d = street()
    d.homes.append(Home(id="Siwama", lat=-1.2903, lon=36.8215, units=24, notes="2 FATs on 2nd floor"))
    autodesign.auto_design(d)
    fats = [n for n in d.nodes if n.kind == "FAT"]
    assert len(fats) == 2 and all(any(s.node == f.id for s in d.splitters) for f in fats)


def test_rerun_numbers_the_same_way():
    d = street()
    d.homes.append(Home(id="Block", lat=-1.2903, lon=36.8215, units=30))
    autodesign.auto_design(d)
    first = sorted(n.id for n in d.nodes if n.status == "proposed")
    autodesign.auto_design(d)
    assert sorted(n.id for n in d.nodes if n.status == "proposed") == first


def test_report_text_has_no_double_full_stop_and_counts_viable_mdus():
    d = sample_design()
    d.homes.append(Home(id="Empty MDU", kind="mdu", units=0, lat=d.homes[0].lat, lon=d.homes[0].lon))
    text = " ".join(p for s in report.sections(d) for p in s.paragraphs)
    assert ".." not in text and "viable MDUs" in text and "no homepasses recorded" in text


def test_long_existing_span_is_used_and_unused_new_poles_are_removed():
    d = street(n=3, step=40.0)
    far = Pole(id="FAR", lat=-1.29, lon=36.82 + 155 / 111_320, status="existing", notes="J-HOOK-1 EXTENSION ARM-1")
    d.poles.append(far)  # 75 m beyond P02, longer than the 65 m max span, but the existing cable already spans it
    d.cables.append(Cable(id="EX 24C", kind="spur", cores=24, status="existing",
                          path=[Vertex(lat=p.lat, lon=p.lon, pole=p.id) for p in d.poles]))
    d.homes.append(Home(id="Block", lat=-1.2902, lon=36.82 + 160 / 111_320, units=24))
    autodesign.auto_design(d)
    used = {v.pole for c in d.cables if c.status == "proposed" for v in c.path}
    assert "FAR" in used and "P02" in used
    assert all(p.id in used or p.status == "existing" for p in d.poles)


@pytest.mark.django_db
def test_every_auto_design_option_is_applied_and_checked(tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    c = client()
    st = c.post("/api/projects", {"sample": True}, content_type="application/json").json()
    r = c.post(f"/api/projects/{st['id']}/auto-design", {"outdoor_fat_ports": 8, "indoor_fat_ports": 16},
               content_type="application/json")
    assert r.status_code == 200
    ports = {n["ports"] for n in r.json()["design"]["nodes"] if n["kind"] == "FAT"}
    assert ports <= {8, 16} and ports
    bad = c.post(f"/api/projects/{st['id']}/auto-design", {"max_span_m": 500}, content_type="application/json")
    assert bad.status_code == 400 and "max_span_m" in bad.json()["detail"]
