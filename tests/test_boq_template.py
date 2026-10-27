"""V Squared BOQ template: catalogue items and rates, template-layout BOQ export, price-list import."""
import io

import pytest
from django.contrib.auth.models import User
from django.test import Client
from openpyxl import load_workbook

from ftth import calc, importers
from ftth.exporters import excel
from ftth.exporters.common import stamp
from ftth.model import DEFAULT_MATERIALS, Design, Material
from ftth.sample import sample_design


def boq_sheet(d):
    wb = load_workbook(io.BytesIO(excel.vsquared_boq_xlsx(d, stamp(d, 1))))
    return wb, wb["BOQ"]


def rows(ws):
    return [[c for c in r if c is not None] for r in ws.iter_rows(values_only=True)]


def test_company_catalogue_has_the_template_items_and_rates():
    tpl = load_workbook(excel.TEMPLATE_BOQ, data_only=True)["Price Lists"]
    want = {str(r[1]).strip(): r[3] for r in tpl.iter_rows(min_row=3, values_only=True) if r[1]}
    have = {"".join(ch for ch in m.description.lower() if ch.isalnum()): m.rate for m in DEFAULT_MATERIALS}
    for desc, rate in want.items():
        assert have["".join(ch for ch in desc.lower() if ch.isalnum())] == rate, desc
    assert {m.code: m.rate for m in DEFAULT_MATERIALS}["EXT-ARM"] == 3500


def test_boq_export_keeps_the_template_layout():
    d = sample_design()
    d.info.name, d.info.company = "Test area FY26", "V Squared"
    wb, ws = boq_sheet(d)
    assert len(ws._images) == 2 and "Price Lists" in wb.sheetnames
    assert ws["C3"].value == "Test area FY26" and ws["F8"].value == "V Squared" and ws["F9"].value == "PLANNED"
    lines = [l for l in calc.bom(d) if l.qty and d.material(l.code).category != "Implementation"]
    body = rows(ws)
    items = [r for r in body if len(r) >= 6 and isinstance(r[-1], str) and r[-1].startswith("=I")]
    assert [r[1] for r in items[4:]] == [l.description for l in lines]  # first 4 are Section A
    total_row = next(r for r in body if r and str(r[0]).startswith("Total Section B"))
    assert total_row[1].startswith("=SUM(M33:")
    names = [r[1] for r in body if r and r[0] in ("Planning Manager ", "Chief Executive Officer")]
    assert names == ["KEVIN MUTURI", "FREDRICK MACHARIA"]
    # the material total the formulas give equals the planner's BOQ total
    assert round(sum(r[3] * r[4] for r in items[4:]), 2) == calc.boq_total([l for l in calc.bom(d)])


def test_implementation_costs_fill_section_a_when_switched_on():
    d = sample_design()
    _, ws = boq_sheet(d)
    assert [ws[f"K{r}"].value for r in range(24, 28)] == [0, 0, 0, 0]
    d.rules.implementation_costs = True
    aerial = sum(c.total_m for c in calc.cable_lengths(d, proposed_only=True) if d.cable(c.cable).install == "aerial")
    _, ws = boq_sheet(d)
    assert ws["K24"].value == aerial and ws["I24"].value == 22


def test_price_list_import_updates_rates_and_adds_items():
    d = Design(materials=[Material(code="J-HOOK", description="Supply of J - Hooks(Small)", unit="PCS", rate=180.0)])
    r = importers.import_price_list(d, excel.TEMPLATE_BOQ.read_bytes())
    codes = {m.code: m for m in d.materials}
    assert r["updated"] == ["J-HOOK: 180 -> 200"] and codes["J-HOOK"].rate == 200
    assert codes["EXT-ARM"].rate == 3500 and codes["INST-AERIAL"].category == "Implementation"
    again = importers.import_price_list(d, excel.TEMPLATE_BOQ.read_bytes())
    assert again["added"] == [] and again["updated"] == []


@pytest.mark.django_db
def test_price_list_upload_endpoint(tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path
    c = Client()
    c.force_login(User.objects.create_user("ann", password="pw-12345-long"))
    pid = c.post("/api/projects", {"name": "A"}, content_type="application/json").json()["id"]
    f = io.BytesIO(excel.TEMPLATE_BOQ.read_bytes())
    f.name = "V SQUARED BOQ.xlsx"
    st = c.post(f"/api/projects/{pid}/prices", {"file": f}).json()
    assert st["prices"]["problems"] == [] and any(m["code"] == "EXT-ARM" for m in st["design"]["materials"])
    assert c.get(f"/api/projects/{pid}/export/boq.xlsx").status_code == 200
