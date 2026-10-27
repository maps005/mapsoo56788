"""The site and documents show the company's logo and details, and return to ArchPlanner once it is deleted."""
import io
import zipfile
from io import BytesIO

import pytest
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from PIL import Image
from pypdf import PdfReader

from ftth.exporters import pdf
from ftth.exporters.common import Brand, stamp
from ftth.sample import sample_design
from planner.branding import DEFAULT_LOGO_FILE, document_brand
from planner.models import Company, Membership, Profile


@pytest.fixture(autouse=True)
def media(tmp_path, settings):
    settings.MEDIA_ROOT = tmp_path


def png_upload():
    b = BytesIO()
    Image.new("RGB", (120, 40), "darkred").save(b, "PNG")
    return SimpleUploadedFile("logo.png", b.getvalue(), content_type="image/png")


def setup_company(u):
    co = Company.objects.create(name="Acme Fibre", legal_name="Acme Fibre Ltd", phone="0700 111222",
                                email="info@acme.example", logo=png_upload())
    Membership.objects.create(user=u, company=co, role="owner")
    Profile.objects.update_or_create(user=u, defaults={"active_company": co})
    return co


def text_of(pdf_bytes: bytes) -> str:
    return " ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(pdf_bytes)).pages)


@pytest.mark.django_db
def test_site_shows_the_company_then_the_default_after_delete():
    u = User.objects.create_user("acme", password="pw-12345-long", is_staff=True, is_superuser=True)
    c = Client()
    c.force_login(u)
    co = setup_company(u)
    for url in ("/", "/projects/", "/admin/"):
        page = c.get(url).content
        assert b"Acme Fibre" in page and f"/company/{co.id}/logo".encode() in page, url
    co.delete()
    for url in ("/", "/projects/", "/admin/"):
        page = c.get(url).content
        assert b"Acme Fibre" not in page and b"archplanner-mark-light.png" in page and b"favicon.png" in page, url


@pytest.mark.django_db
def test_documents_carry_the_company_then_the_default():
    u = User.objects.create_user("acme")
    co = setup_company(u)
    b = document_brand(co)
    assert b.name == "Acme Fibre Ltd" and "0700 111222" in b.details and b.logo and b.logo.is_file()
    d = sample_design()
    out = text_of(pdf.bom_pdf(d, stamp(d, 1, b)))
    assert "Acme Fibre Ltd" in out and "info@acme.example" in out
    co.delete()
    default = document_brand(None)
    assert default.name == "ArchPlanner" and default.logo == DEFAULT_LOGO_FILE
    out = text_of(pdf.bom_pdf(d, stamp(d, 1, default)))
    assert "ArchPlanner" in out and "Acme" not in out


@pytest.mark.django_db
def test_export_uses_the_projects_company():
    u = User.objects.create_user("acme", password="pw-12345-long")
    setup_company(u)
    c = Client()
    c.force_login(u)
    pid = c.post("/api/projects", {"sample": True}, content_type="application/json").json()["id"]
    r = c.get(f"/api/projects/{pid}/export/bom.pdf")
    assert r.status_code == 200 and "Acme Fibre Ltd" in text_of(b"".join(r.streaming_content) if r.streaming else r.content)


def test_every_document_builds_with_and_without_a_logo(tmp_path):
    from ftth.exporters import package
    d = sample_design()
    for brand in (Brand(), Brand(name="X Ltd", details="Nairobi", logo=DEFAULT_LOGO_FILE)):
        z = zipfile.ZipFile(io.BytesIO(package.build_package(d, stamp(d, 1, brand), tmp_path)))
        assert len(z.namelist()) > 10


def test_survey_report_is_a_full_document():
    from docx import Document

    from ftth.exporters import word
    d = sample_design()
    st = stamp(d, 2, Brand(name="Acme Fibre Ltd", logo=DEFAULT_LOGO_FILE))
    text = text_of(pdf.report_pdf(d, st, None))
    for part in ("FTTH SURVEY AND DESIGN REPORT", "Document control", "Sign-off", "Contents", "Executive Summary",
                 "3.1 Network elements", "6. Bill of Quantities Summary", "7. Quality Assurance", "Outstanding actions"):
        assert part in text, part
    docx = Document(io.BytesIO(word.report_docx(d, st, None)))
    heads = [p.text for p in docx.paragraphs if p.style.name.startswith("Heading")]
    assert "Executive Summary" in heads and "9. Conclusion and Recommendations" in heads
    assert "Acme Fibre Ltd" in docx.sections[0].header.paragraphs[0].text
