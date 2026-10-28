"""Excel deliverables: pole summary, BOM, BOQ (live formulas) and full project data."""
from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .. import calc, splicing
from ..model import Design
from .common import Stamp, clean

HEAD = PatternFill("solid", fgColor="1F4E79")
SUB = PatternFill("solid", fgColor="D9E2F3")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sheet(wb: Workbook, title: str, st: Stamp, heading: str, cols: list[tuple[str, int]], first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws["A1"] = clean(heading)
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws["A2"] = clean(st.line)
    ws["A2"].font = Font(italic=True, size=9, color="595959")
    for i, (name, width) in enumerate(cols, start=1):
        c = ws.cell(row=4, column=i, value=name)
        c.fill, c.font, c.border = HEAD, Font(bold=True, color="FFFFFF"), BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A5"
    return ws


class Formula(str):
    """A formula written by this tool (other text that starts with '=' is stored as plain text)."""


def _row(ws, r: int, values: list, bold=False, fill=None, fmt: dict[int, str] | None = None):
    for i, v in enumerate(values, start=1):
        if isinstance(v, str) and not isinstance(v, Formula):
            v = clean(v)
        c = ws.cell(row=r, column=i, value=v)
        if isinstance(v, str) and not isinstance(v, Formula) and v[:1] in ("=", "+", "-", "@"):
            c.data_type = "s"
        c.border = BOX
        if bold:
            c.font = Font(bold=True)
        if fill:
            c.fill = fill
        if fmt and i in fmt:
            c.number_format = fmt[i]


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sector_schedule_xlsx(d: Design, st: Stamp) -> bytes:
    """Other sectors: the support schedule (manholes, marker posts, poles, columns) with the sector's hardware,
    equipment and recorded details."""
    from .. import sectors
    sec = sectors.get(d.info.sector)
    cols = sectors.schedule_columns(d.info.sector)
    fields = sec["details"].get("pole", [])
    wb = Workbook()
    heads = [(f"{sec['support']} ID", 12), ("Latitude", 11), ("Longitude", 11), ("Status", 10),
             (f"{sec['line']}s", 8), (f"{sec['line']} IDs", 18)] + [(h, 11) for _, h in cols] + \
        [(f["label"] + (f" ({f['unit']})" if f["unit"] else ""), 12) for f in fields] + [("Notes", 40)]
    ws = _sheet(wb, f"{sec['support']} Schedule"[:31], st, f"{sec['support']} Schedule - {d.info.name}", heads, True)
    r = 5
    for p in calc.pole_summary(d):
        det = (d.pole(p.pole).details if d.pole(p.pole) else {})
        _row(ws, r, [p.pole, p.lat, p.lon, p.status, len(p.cables), ", ".join(p.cables)] +
             [getattr(p, k) for k, _ in cols] + [det.get(f["key"]) for f in fields] + ["; ".join(p.notes)],
             fmt={2: "0.000000", 3: "0.000000"})
        r += 1
    first, last = 7, 6 + len(cols)
    _row(ws, r, ["Total", None, None, None, Formula(f"=SUM(E5:E{r - 1})"), None] +
         [Formula(f"=SUM({get_column_letter(c)}5:{get_column_letter(c)}{r - 1})") for c in range(first, last + 1)] +
         [None] * (len(fields) + 1), bold=True, fill=SUB)
    return _save(wb)


def pole_summary_xlsx(d: Design, st: Stamp) -> bytes:
    from .. import sectors
    if not sectors.fibre_tools(d):
        return _sector_schedule_xlsx(d, st)
    wb = Workbook()
    ws = _sheet(wb, "Pole Summary", st, f"Pole Summary - {d.info.name}", [
        ("Pole ID", 10), ("Latitude", 11), ("Longitude", 11), ("Status", 10), ("Cables", 8), ("Cable IDs", 18),
        ("Tension Clamps", 9), ("UPB", 7), ("J-Hooks", 8), ("Extension Arm", 9), ("JB", 6), ("Enclosure", 9),
        ("FAT", 6), ("Notes", 40)], True)
    rows = calc.pole_summary(d)
    r = 5
    for p in rows:
        _row(ws, r, [p.pole, p.lat, p.lon, p.status, len(p.cables), ", ".join(p.cables), p.tension, p.upb,
                     p.jhooks, p.ext_arms, p.jb, p.enclosures, p.fats, "; ".join(p.notes)],
             fmt={2: "0.000000", 3: "0.000000"})
        r += 1
    _row(ws, r, ["Total", None, None, None] + [Formula(f"=SUM({get_column_letter(c)}5:{get_column_letter(c)}{r - 1})")
                                               if c != 6 else None for c in range(5, 14)] + [None],
         bold=True, fill=SUB)
    rr = d.rules
    ws.cell(row=r + 2, column=1, value="Rules applied:").font = Font(bold=True)
    ws.cell(row=r + 3, column=1, value=(
        f"Tension clamps: {rr.tension_per_cable_end} per cable end, {rr.tension_per_cable_corner} per cable at corners "
        f"(deflection >= {rr.corner_angle_deg:g} deg) and at anchor poles (after {rr.max_suspension_poles} J-hook poles "
        f"in a row). J-hook: {rr.jhook_per_straight_cable} per cable at other straight poles. UPB: 1 per "
        f"{rr.upb_tension_per_bracket} tension clamps. Extension arm on {', '.join(rr.extension_arm_pole_types) or '-'} "
        f"poles. Existing cables {'included' if rr.accessories_on_existing_cables else 'excluded'}."))
    return _save(wb)


def bom_xlsx(d: Design, st: Stamp) -> bytes:
    wb = Workbook()
    ws = _sheet(wb, "BOM", st, f"Bill of Materials - {d.info.name}", [
        ("Item", 6), ("Code", 16), ("Description", 34), ("Unit", 6), ("Calculated", 11), ("Adjustment", 11),
        ("Quantity", 11), ("Basis", 30)], True)
    r = 5
    for i, l in enumerate(calc.bom(d), start=1):
        _row(ws, r, [i, l.code, l.description, l.unit, l.calculated, l.adjustment, Formula(f"=E{r}+F{r}"), l.source])
        r += 1
    _adjustments_sheet(wb, d, st)
    return _save(wb)


def boq_xlsx(d: Design, st: Stamp) -> bytes:
    wb = Workbook()
    cur = d.info.currency
    ws = _sheet(wb, "BOQ", st, f"Bill of Quantities - {d.info.name}", [
        ("Item", 6), ("Code", 16), ("Description", 34), ("Unit", 6), ("Quantity", 11),
        (f"Unit Rate ({cur})", 13), (f"Total ({cur})", 14)], True)
    r = 5
    lines = calc.bom(d)
    for i, l in enumerate(lines, start=1):
        _row(ws, r, [i, l.code, l.description, l.unit, l.qty, l.rate, Formula(f"=IF(F{r}=\"\",0,E{r}*F{r})")],
             fmt={6: "#,##0.00", 7: "#,##0.00"})
        r += 1
    _row(ws, r, [None, None, "TOTAL", None, None, None, Formula(f"=SUM(G5:G{r - 1})")], bold=True, fill=SUB,
         fmt={7: "#,##0.00"})
    missing = [l.code for l in lines if l.rate is None]
    if missing:
        ws.cell(row=r + 2, column=1, value=f"No rate in catalogue for: {', '.join(missing)}").font = Font(color="C00000")
    return _save(wb)


def _adjustments_sheet(wb, d: Design, st: Stamp):
    ws = _sheet(wb, "Adjustments", st, "Authorised manual adjustments", [
        ("Code", 16), ("Quantity change", 14), ("Reason", 40), ("By", 16)])
    for i, a in enumerate(d.adjustments, start=5):
        _row(ws, i, [a.code, a.qty, a.reason, a.by])


def _path_token(v) -> str:
    """Cable point as written in the Cables sheet: pole, equipment, 'equipment @ pole', or 'pt' for a bend."""
    if v.node and v.pole:
        return f"{v.node} @ {v.pole}"
    return v.node or v.pole or "pt"


def project_xlsx(d: Design, st: Stamp) -> bytes:
    wb = Workbook()
    specs = [
        ("Poles", ["id", "lat", "lon", "type", "height_m", "condition", "status", "notes"], d.poles),
        ("Homes", ["id", "lat", "lon", "kind", "units", "name", "address", "fat", "access_notes", "notes"], d.homes),
        ("Equipment", ["id", "kind", "lat", "lon", "pole", "status", "ports", "notes"], d.nodes),
        ("Splitters", ["id", "node", "stage", "outputs"], d.splitters),
    ]
    first = True
    for title, fields, items in specs:
        ws = _sheet(wb, title, st, f"{title} - {d.info.name}", [(f, 14) for f in fields], first)
        first = False
        for r, it in enumerate(items, start=5):
            _row(ws, r, [getattr(it, f) for f in fields])
    lengths = {cl.cable: cl for cl in calc.cable_lengths(d)}
    ws = _sheet(wb, "Cables", st, "Cable schedule", [(h, 13) for h in (
        "id", "kind", "cores", "from", "to", "install", "status", "route m", "extra m", "slack m", "riser m",
        "total m", "crossings", "path", "path coordinates")])
    for r, c in enumerate(d.cables, start=5):
        cl = lengths.get(c.id)
        metres = [round(x, 1) for x in (cl.route_m, cl.extra_m, cl.slack_m, cl.riser_m, cl.total_m)] if cl else [None] * 5
        _row(ws, r, [c.id, c.kind, c.cores, c.from_node, c.to_node, c.install, c.status, *metres,
                     cl.crossings if cl else None, " > ".join(_path_token(v) for v in c.path),
                     "; ".join(f"{v.lat:.7f},{v.lon:.7f}" for v in c.path)])
    ws = _sheet(wb, "Fibre Cores", st, "Fibre core allocation", [
        ("Cable", 10), ("Core", 6), ("Tube", 6), ("Tube colour", 11), ("Fibre colour", 11), ("Status", 10), ("Connections", 60)])
    r = 5
    for cab in splicing.core_table(d):
        for c in cab["core_list"]:
            _row(ws, r, [cab["cable"], c["n"], c["tube"], c["tube_colour"], c["fibre_colour"], c["status"], " | ".join(c["uses"])])
            r += 1
    ws = _sheet(wb, "Connections", st, "Splice / connection records", [
        ("ID", 8), ("Node", 10), ("A", 34), ("B", 34), ("Auto", 6), ("Note", 30)])
    for r, c in enumerate(d.connections, start=5):
        _row(ws, r, [c.id, c.node, splicing.end_label(d, c.a), splicing.end_label(d, c.b), "yes" if c.auto else "", c.note])
    ws = _sheet(wb, "Materials", st, "Material catalogue", [
        ("Code", 16), ("Description", 34), ("Unit", 6), ("Rate", 10), ("Category", 12)])
    for r, m in enumerate(d.materials, start=5):
        _row(ws, r, [m.code, m.description, m.unit, m.rate, m.category])
    return _save(wb)


def project_csvs(d: Design) -> dict[str, str]:
    import csv
    out = {}
    for name, fields, items in (
        ("poles", ["id", "lat", "lon", "type", "height_m", "condition", "status", "notes"], d.poles),
        ("homes", ["id", "lat", "lon", "kind", "units", "name", "address", "fat", "access_notes", "notes"], d.homes),
        ("equipment", ["id", "kind", "lat", "lon", "pole", "status", "ports", "notes"], d.nodes),
    ):
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(fields)
        for it in items:
            w.writerow([getattr(it, f) if getattr(it, f) is not None else "" for f in fields])
        out[f"{name}.csv"] = buf.getvalue()
    return out


def home_list_xlsx(d: Design, st: Stamp) -> bytes:
    """Home list in the V Squared layout: MDUs, SDUs and commercial units with viable homepasses."""
    from openpyxl.styles import Font as F
    wb = Workbook()
    ws = wb.active
    ws.title = clean(d.info.cluster or d.info.name or "Home list")[:31] or "Home list"
    headers = ["REGION", "ZONE", "CLUSTER", "GREEN/BROWN FIELD", "NAME", "NATURE OF BUILDING", "VIABLE HP",
               "COMMERCIAL", "LONGITUDE", "LATITUDE", "ACQUIRED", "SERVING FAT"]
    widths = [12, 12, 12, 16, 34, 18, 10, 12, 13, 13, 11, 18]
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill, c.font, c.border = HEAD, F(bold=True, color="FFFFFF"), BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    sections = [("MULTI DWELLING UNITS (MDUs)", "mdu", "MDU"), ("SINGLE DWELLING UNITS (SDUs)", "home", "SDU"),
                ("COMMERCIAL UNITS", "business", "COMMERCIAL")]
    info = d.info
    r, totals = 2, {}
    for title, kind, label in sections:
        homes = [h for h in d.homes if h.kind == kind]
        ws.cell(row=r, column=1, value=title).font = F(bold=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
        ws.cell(row=r, column=1).fill = SUB
        r += 1
        for h in homes:
            field = {"green": "GREEN FIELD", "brown": "BROWN FIELD"}.get(h.field_type, "")
            _row(ws, r, [info.region.upper(), info.zone.upper(), info.cluster.upper(), field,
                         (h.name or h.id).upper(), label, h.units or None, None, h.lon, h.lat,
                         "ACQUIRED" if h.acquired else None, ", ".join(h.fats) or None])
            r += 1
        totals[label] = sum(h.units for h in homes)
        _row(ws, r, [None] * 5 + ["TOTAL", totals[label]] + [None] * 5, bold=True)
        r += 2
    ws.cell(row=r, column=7, value="BUILDING TYPE").font = F(bold=True)
    ws.cell(row=r, column=8, value="QUANTITY").font = F(bold=True)
    for i, (label, name) in enumerate((("MDU", "MDU"), ("SDU", "SDU"), ("COMMERCIAL", "BUSINESS UNITS")), start=1):
        ws.cell(row=r + i, column=7, value=name)
        ws.cell(row=r + i, column=8, value=totals.get(label, 0))
    ws.cell(row=r + 4, column=7, value="TOTAL").font = F(bold=True)
    ws.cell(row=r + 4, column=8, value=sum(totals.values())).font = F(bold=True)
    ws.cell(row=r + 6, column=1, value=clean(st.line)).font = F(italic=True, size=9, color="595959")
    ws.freeze_panes = "A2"
    return _save(wb)


# ---------------------------------------------------------------------------
# Survey upload template
# ---------------------------------------------------------------------------
TEMPLATE = {  # sheet -> [(column, width, required, example, meaning)]
    "Poles": [("ID", 10, "yes", "P001", "Unique pole number"),
              ("Latitude", 12, "yes", -1.292066, "Decimal degrees (WGS84)"),
              ("Longitude", 12, "yes", 36.821946, "Decimal degrees (WGS84)"),
              ("Type", 10, "", "wood", "wood, concrete, steel"),
              ("Height", 8, "", 9, "Metres"),
              ("Condition", 10, "", "good", "good, leaning, rotten, ..."),
              ("Status", 10, "", "existing", "existing or proposed"),
              ("Notes", 30, "", "", "")],
    "Homes": [("ID", 10, "yes", "H001", "Unique premise number"),
              ("Latitude", 12, "yes", -1.292100, ""), ("Longitude", 12, "yes", 36.822000, ""),
              ("Name", 24, "", "Riverside Apartments", "Building or customer name"),
              ("Nature of building", 14, "", "MDU", "MDU, SDU / home, business"),
              ("Viable HP", 9, "yes", 24, "Homepasses (units) in the building"),
              ("Address", 20, "", "", ""), ("Acquired", 9, "", "", "yes when wayleave / access is acquired"),
              ("Green/Brown field", 12, "", "brown", "green or brown"),
              ("FAT", 10, "", "", "Serving FAT, if already decided"),
              ("Access", 20, "", "", "Access notes / restrictions"), ("Notes", 24, "", "", "")],
    "Equipment": [("ID", 12, "yes", "Ex Closure 01", "Unique equipment name"),
                  ("Kind", 12, "yes", "ENCLOSURE", "POP, ENCLOSURE, PR_ENCLOSURE, JB, PON, FAT, SLACK"),
                  ("Latitude", 12, "yes", -1.291900, ""), ("Longitude", 12, "yes", 36.821500, ""),
                  ("Pole", 10, "", "P001", "Pole it is mounted on"),
                  ("Status", 10, "", "existing", "existing or proposed"),
                  ("Ports", 7, "", "", "FAT ports (8, 16, 24)"), ("Notes", 24, "", "", "")],
    "Splitters": [("ID", 10, "", "SP01", "Leave blank to number automatically"),
                  ("Node", 12, "yes", "FAT-01", "Equipment ID the splitter is in"),
                  ("Stage", 7, "", 2, "1 = primary (PON), 2 = FAT"),
                  ("Outputs", 8, "", "1x8", "1x4, 1x8, 1x16 ...")],
    "Cables": [("ID", 10, "", "CB01", "Leave blank to number automatically"),
               ("Kind", 9, "", "main", "main (feeder), spur (distribution), drop"),
               ("Cores", 7, "yes", 48, "Fibre count"),
               ("From", 12, "", "Ex Closure 01", "Start equipment ID"),
               ("To", 12, "", "FAT-01", "End equipment ID"),
               ("Path", 40, "yes", "Ex Closure 01 > P001 > P002 > FAT-01",
                "Pole / equipment IDs in route order, separated by > or commas"),
               ("Install", 11, "", "aerial", "aerial or underground"),
               ("Status", 10, "", "proposed", "existing or proposed"), ("Notes", 24, "", "", "")],
}
TEMPLATE_LISTS = {("*", "Status"): "existing,proposed", ("*", "Install"): "aerial,underground",
                  ("Equipment", "Kind"): "POP,ENCLOSURE,PR_ENCLOSURE,JB,PON,FAT,SLACK", ("Cables", "Kind"): "main,spur,drop"}


def survey_template_xlsx() -> bytes:
    """Blank workbook for survey data: one sheet per kind of record, plus instructions with examples."""
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    ws["A1"] = "FTTH survey data upload template"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    lines = ["Fill in the sheets you need (any can stay empty) and upload the whole file under Survey data > Import",
             "with 'Rows are' set to auto. Sheets are read in this order: Poles, Equipment, Homes, Splitters, Cables,",
             "so cables can refer to pole and equipment IDs in the same file. A row with an existing ID updates that",
             "record. Rows with missing coordinates are kept and flagged for correction. Keep the sheet names.",
             "Dark blue headings are required columns."]
    for i, t in enumerate(lines, start=2):
        ws.cell(row=i, column=1, value=t)
    r = len(lines) + 3
    for i, (h, w) in enumerate((("Sheet", 12), ("Column", 18), ("Required", 9), ("Example", 34), ("Meaning", 56)), 1):
        c = ws.cell(row=r, column=i, value=h)
        c.fill, c.font, c.border = HEAD, Font(bold=True, color="FFFFFF"), BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    for sheet, cols in TEMPLATE.items():
        for name, _w, req, ex, meaning in cols:
            r += 1
            _row(ws, r, [sheet, name, req, ex, meaning])
    for sheet, cols in TEMPLATE.items():
        sh = wb.create_sheet(sheet)
        for i, (name, width, req, _ex, _m) in enumerate(cols, start=1):
            c = sh.cell(row=1, column=i, value=name)
            c.fill, c.font, c.border = (HEAD if req else SUB), Font(bold=True, color="FFFFFF" if req else "1F4E79"), BOX
            sh.column_dimensions[get_column_letter(i)].width = width
            choices = TEMPLATE_LISTS.get((sheet, name)) or TEMPLATE_LISTS.get(("*", name))
            if choices:
                dv = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
                sh.add_data_validation(dv)
                dv.add(f"{get_column_letter(i)}2:{get_column_letter(i)}5000")
        sh.freeze_panes = "A2"
    return _save(wb)


# ---------------------------------------------------------------------------
# Report workbooks (the same content as the PDF reports)
# ---------------------------------------------------------------------------
def splicing_xlsx(d: Design, st: Stamp) -> bytes:
    wb = Workbook()
    ws = _sheet(wb, "Splitters", st, f"Splitters - {d.info.name}", [
        ("Node", 14), ("Node type", 12), ("Splitter", 12), ("Stage", 7), ("Ratio", 8), ("Fed from (OLT)", 30)], True)
    kinds = {n.id: n.kind for n in d.nodes}
    for r, s in enumerate(d.splitters, start=5):
        tr = splicing.trace_upstream(d, s.id)
        _row(ws, r, [s.node, kinds.get(s.node, ""), s.id, s.stage, s.ratio,
                     tr.get("olt") if tr["ok"] else f"NOT FED: {tr.get('reason', '')}"])
    ws = _sheet(wb, "Splice Records", st, "Splice / connection records by node", [
        ("Node", 14), ("ID", 8), ("Incoming", 40), ("Outgoing", 40), ("Source", 8), ("Note", 24)])
    notes = {c.id: c.note for c in d.connections}
    r = 5
    for n in d.nodes:
        for c in splicing.node_splicing(d, n.id)["connections"]:
            _row(ws, r, [n.id, c["id"], c["a_label"], c["b_label"], "auto" if c["auto"] else "manual", notes.get(c["id"], "")])
            r += 1
    cores = splicing.core_table(d)
    ws = _sheet(wb, "Core Usage", st, "Core usage by cable", [
        ("Cable", 12), ("Type", 8), ("Cores", 7), ("Used", 7), ("Reserved", 9), ("Spare", 7)])
    for r, c in enumerate(cores, start=5):
        _row(ws, r, [c["cable"], c["kind"], c["cores"], c["used"], c["reserved"], c["spare"]])
    ws = _sheet(wb, "Fibre Cores", st, "Fibre core allocation", [
        ("Cable", 10), ("Core", 6), ("Tube", 6), ("Tube colour", 11), ("Fibre colour", 11), ("Status", 10), ("Connections", 60)])
    r = 5
    for cab in cores:
        for c in cab["core_list"]:
            _row(ws, r, [cab["cable"], c["n"], c["tube"], c["tube_colour"], c["fibre_colour"], c["status"], " | ".join(c["uses"])])
            r += 1
    ws = _sheet(wb, "FAT Traces", st, "FAT to OLT traces", [("FAT", 14), ("Splitter", 12), ("Result", 30), ("Path", 90)])
    r = 5
    for s in d.splitters:
        if s.stage != 2:
            continue
        tr = splicing.trace_upstream(d, s.id)
        _row(ws, r, [s.node, s.id, tr.get("olt") if tr["ok"] else f"NOT FED: {tr['reason']}",
                     " | ".join(f"{x['node']}: {x['from']} -> {x['to']}" for x in tr["steps"])])
        r += 1
    return _save(wb)


def report_xlsx(d: Design, st: Stamp) -> bytes:
    from . import report
    wb = Workbook()
    ws = wb.active
    ws.title = "Survey Report"
    from .. import sectors
    ws["A1"] = sectors.get(d.info.sector)["report_title"].upper()
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws["A2"] = clean(f"{d.info.name} - {d.info.area}")
    ws["A3"] = clean(f"Prepared by: {d.info.prepared_by or '-'}   Date: {st.date}   Revision {st.revision} ({st.status})")
    ws["A4"] = clean(st.line)
    ws["A4"].font = Font(italic=True, size=9, color="595959")
    for col, w in zip("ABCDEFGH", (24, 18, 14, 14, 14, 14, 14, 16)):
        ws.column_dimensions[col].width = w
    r = 6
    for sec in report.sections(d):
        ws.cell(row=r, column=1, value=sec.title).font = Font(bold=True, size=12, color="1F4E79")
        r += 1
        for p in sec.paragraphs:
            c = ws.cell(row=r, column=1, value=p)
            c.data_type = "s"
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
            ws.row_dimensions[r].height = 15 * (len(p) // 110 + 1)
            r += 1
        for headers, rows in sec.tables:
            for i, h in enumerate(headers, start=1):
                c = ws.cell(row=r, column=i, value=h)
                c.fill, c.font, c.border = HEAD, Font(bold=True, color="FFFFFF"), BOX
            r += 1
            for row in rows:
                _row(ws, r, list(row))
                r += 1
        r += 1
    if d.photos:
        ph = _sheet(wb, "Photos", st, "Survey photographs (the images are in the PDF and Word report)", [
            ("File", 18), ("Caption", 40), ("Feature", 14)])
        for i, p in enumerate(d.photos, start=5):
            _row(ph, i, [p.file, p.caption, p.feature])
    return _save(wb)


def qa_rows(d: Design, st: Stamp) -> tuple[list[list], list[dict]]:
    """QA summary rows and the validation issues (errors first), shared by the Excel and PDF QA reports."""
    from .. import validate
    issues = validate.validate(d)
    cnt = validate.counts(issues)
    i, s = d.info, calc.summary(d)
    rows = [["Survey dataset", i.survey_status], ["Survey approved by", i.survey_approved_by or "-"],
            ["Survey approved at", i.survey_approved_at or "-"], ["Design status", i.design_status],
            ["Revision", st.revision], ["Data hash", st.hash], ["Planned by", i.planned_by or "-"],
            ["Checked by", i.checked_by or "-"], ["Approved by", i.approved_by or "-"],
            ["Validation errors", cnt["errors"]], ["Validation warnings", cnt["warnings"]], ["Notes (information only)", cnt["notes"]],
            ["Homepasses", s["homepasses"]], ["BOQ total", f"{s['currency']} {s['boq_total']:,.2f}"],
            ["Ready for approval", "yes" if cnt["errors"] == 0 else "no - fix the errors first"]]
    order = {"error": 0, "warning": 1}
    return rows, sorted(issues, key=lambda x: (order.get(x["severity"], 2), x["code"]))


def qa_xlsx(d: Design, st: Stamp) -> bytes:
    rows, issues = qa_rows(d, st)
    wb = Workbook()
    ws = _sheet(wb, "QA Summary", st, f"QA and approval - {d.info.name}", [("Item", 26), ("Detail", 50)], True)
    for r, row in enumerate(rows, start=5):
        _row(ws, r, row)
    ws = _sheet(wb, "Validation", st, "Validation issues", [("Severity", 10), ("Check", 22), ("Feature", 16), ("Issue", 90)])
    for r, x in enumerate(issues, start=5):
        _row(ws, r, [x["severity"], x["code"], x.get("ref") or "", x["message"]])
        if x["severity"] == "error":
            ws.cell(row=r, column=1).font = Font(bold=True, color="C00000")
    return _save(wb)


# ---------------------------------------------------------------------------
# BOQ in the V Squared company template (logo, overview, Section A / B, approval block)
# ---------------------------------------------------------------------------
from copy import copy
from pathlib import Path

TEMPLATE_BOQ = Path(__file__).resolve().parents[1] / "templates" / "vsquared_boq.xlsx"
_ITEM_COLS = "CDEFGHIJKLMN"
_ITEM_MERGES = ("C{r}:D{r}", "E{r}:F{r}", "G{r}:H{r}", "I{r}:J{r}", "K{r}:L{r}", "M{r}:N{r}")
_SECTION_A_ROWS = range(24, 28)  # template rows: aerial installation, trunking, trench, building cabling


def _row_style(ws, r):
    return {c: copy(ws[f"{c}{r}"]._style) for c in _ITEM_COLS}, ws.row_dimensions[r].height


def _apply_style(ws, r, style):
    cells, height = style
    for c, st in cells.items():
        ws[f"{c}{r}"]._style = copy(st)
    ws.row_dimensions[r].height = height


def _norm_desc(s) -> str:
    return "".join(ch for ch in str(s or "").lower() if ch.isalnum())


def vsquared_boq_xlsx(d: Design, st: Stamp) -> bytes:
    """The BOQ in the company template: Section A implementation and Section B materials with live formulas."""
    if not TEMPLATE_BOQ.exists():
        return boq_xlsx(d, st)
    wb = load_workbook(TEMPLATE_BOQ)
    ws = wb["BOQ"]
    i = d.info
    status = {"draft": "PLANNED", "in_review": "IN REVIEW", "approved": "APPROVED"}[i.design_status]
    ws["C3"] = clean(i.name)
    ws["F6"] = clean(i.name)
    ws["F7"] = "FTTH"
    ws["F8"] = clean(i.company)
    ws["F9"] = status
    day = "/".join(reversed(st.date.split("-")))
    ws["F10"], ws["F11"] = day, day
    lines = calc.bom(d)
    impl_codes = {m.code for m in d.materials if m.category == "Implementation"}
    impl = {l.code: l for l in lines if l.code in impl_codes}
    mats = [l for l in lines if l.code not in impl_codes and l.qty]
    # Section A: the template's four rows, matched to the catalogue by description
    by_desc = {_norm_desc(m.description): m for m in d.materials}
    for r in _SECTION_A_ROWS:
        m = by_desc.get(_norm_desc(ws[f"E{r}"].value))
        line = impl.get(m.code) if m else None
        if m:
            ws[f"I{r}"] = m.rate
        ws[f"K{r}"] = line.qty if line else 0
    # Section B: rebuild the item rows from the template's own row styles
    item, total = _row_style(ws, 33), _row_style(ws, 53)
    tail = [(r, _row_style(ws, r), [ws[f"{c}{r}"].value for c in _ITEM_COLS]) for r in range(54, 63)]
    for m in [m for m in ws.merged_cells.ranges if m.min_row >= 33]:
        ws.unmerge_cells(str(m))
    for r in range(33, ws.max_row + 1):  # clear values and formatting below Section B's heading
        for c in _ITEM_COLS:
            ws[f"{c}{r}"].value = None
            ws[f"{c}{r}"].style = "Normal"
        ws.row_dimensions[r].height = None
    first = r = 33
    for l in mats:
        _apply_style(ws, r, item)
        for m in _ITEM_MERGES:
            ws.merge_cells(m.format(r=r))
        mat = d.material(l.code)
        ws[f"C{r}"] = clean(mat.category if mat else "")
        ws[f"E{r}"] = clean(l.description)
        ws[f"G{r}"] = clean(l.unit)
        ws[f"I{r}"] = l.rate if l.rate is not None else None
        ws[f"K{r}"] = l.qty
        ws[f"M{r}"] = f"=I{r}*K{r}"
        r += 1
    tr = r
    _apply_style(ws, tr, total)
    ws.merge_cells(f"C{tr}:L{tr}")
    ws.merge_cells(f"M{tr}:N{tr}")
    ws[f"C{tr}"] = "Total Section B Material Costs : "
    ws[f"M{tr}"] = f"=SUM(M{first}:M{tr - 1})" if tr > first else 0
    ws["I16"] = f"=M{tr}"
    # approval block (titles and names from the template) below the total
    shift = tr - 53
    for r0, style, values in tail:
        r1 = r0 + shift
        _apply_style(ws, r1, style)
        for c, v in zip(_ITEM_COLS, values):
            ws[f"{c}{r1}"] = v
        if r0 == 56:
            ws.merge_cells(f"C{r1}:N{r1}")
        elif r0 >= 57:
            for m in (f"C{r1}:E{r1}", f"G{r1}:K{r1}", f"L{r1}:N{r1}"):
                ws.merge_cells(m)
    note = 62 + shift + 2
    ws[f"C{note}"] = clean(st.line)
    ws[f"C{note}"].font = Font(italic=True, size=8, color="595959")
    missing = [l.code for l in mats if l.rate is None]
    if missing:
        ws[f"C{note + 1}"] = f"No rate in catalogue for: {', '.join(missing)}"
        ws[f"C{note + 1}"].font = Font(bold=True, color="C00000")
    ws.print_area = f"A1:P{note + 1}"
    wb.calculation.fullCalcOnLoad = True  # totals are formulas: recalculate when opened
    # the price list sheet shows the project's catalogue
    if "Price Lists" in wb.sheetnames:
        pl = wb["Price Lists"]
        style = [copy(pl.cell(row=3, column=k)._style) for k in range(1, 5)]
        for row in pl.iter_rows(min_row=3):
            for cell in row:
                cell.value = None
        for n, m in enumerate([m for m in d.materials if m.category != "Implementation"], start=3):
            for k, v in enumerate((m.category, m.description, m.unit, m.rate), start=1):
                cell = pl.cell(row=n, column=k, value=clean(v) if isinstance(v, str) else v)
                cell._style = copy(style[k - 1])
    return _save(wb)
