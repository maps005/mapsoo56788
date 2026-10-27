"""Editable Word (.docx) survey and design report (same content as the PDF report)."""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..model import Design
from . import report
from .common import Stamp, clean, logo_size, photo_path

NAVY, BLUE, GOLD, GREY = RGBColor(0x0B, 0x25, 0x45), RGBColor(0x15, 0x57, 0xB0), RGBColor(0xB8, 0x86, 0x2B), \
    RGBColor(0x4A, 0x5B, 0x73)


def _shade(cell, hex_fill: str):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc.append(shd)


def _cell(cell, text, bold=False, colour=None, size=8.5, right=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if right:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("" if text is None else str(text))
    r.bold, r.font.size = bold, Pt(size)
    if colour is not None:
        r.font.color.rgb = colour


def _table(doc, headers, rows, total=False):
    tb = doc.add_table(rows=1, cols=len(headers))
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.LEFT
    numeric = [all(str(r[j]).replace(",", "").replace(".", "").replace(" m", "").replace("KES", "").strip().isdigit()
                   or not str(r[j]).strip() for r in rows) and any(str(r[j]).strip() for r in rows)
               for j in range(len(headers))]
    for j, h in enumerate(headers):
        _cell(tb.rows[0].cells[j], h, bold=True, colour=RGBColor(0xFF, 0xFF, 0xFF))
        _shade(tb.rows[0].cells[j], "0B2545")
    for n, row in enumerate(rows):
        cells = tb.add_row().cells
        last = total and n == len(rows) - 1
        for j, v in enumerate(row):
            _cell(cells[j], v, bold=last, right=numeric[j])
            if last:
                _shade(cells[j], "D9E2F3")
            elif n % 2:
                _shade(cells[j], "F2F5FA")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _kv(doc, rows):
    tb = doc.add_table(rows=0, cols=2)
    tb.style = "Table Grid"
    for k, v in rows:
        cells = tb.add_row().cells
        _cell(cells[0], k, bold=True)
        _shade(cells[0], "EEF3FA")
        _cell(cells[1], v)
        cells[0].width, cells[1].width = Cm(5.5), Cm(11.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _heading(doc, text, level):
    h = doc.add_heading(text, level)
    for r in h.runs:
        r.font.color.rgb = NAVY if level == 1 else BLUE
    return h


def report_docx(d: Design, st: Stamp, photo_dir: Path | None) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.different_first_page_header_footer = True  # the cover page has no running header
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10.5)

    sec0 = doc.sections[0]
    hdr = sec0.header.paragraphs[0]  # issuer logo and name on every page after the cover
    size = logo_size(st.brand.logo, 3.6, 1.1)
    if size:
        try:
            hdr.add_run().add_picture(str(st.brand.logo), width=Cm(size[0]), height=Cm(size[1]))
            hdr.add_run("   ")
        except (OSError, ValueError):
            pass
    hr = hdr.add_run(clean(st.brand.name))
    hr.bold, hr.font.size, hr.font.color.rgb = True, Pt(10), NAVY
    if st.brand.details:
        dr = sec0.header.add_paragraph().add_run(clean(st.brand.details))
        dr.font.size, dr.font.color.rgb = Pt(8), GREY
    ft = sec0.footer.paragraphs[0].add_run(clean(f"{d.info.name}  |  FTTH Survey and Design Report  |  "
                                                 f"Rev {st.revision} ({st.status.replace('_', ' ')})  |  {st.date}"))
    ft.font.size, ft.font.color.rgb = Pt(8), GREY

    # cover
    size = logo_size(st.brand.logo, 6.5, 2.8)
    if size:
        try:
            doc.add_picture(str(st.brand.logo), width=Cm(size[0]), height=Cm(size[1]))
        except (OSError, ValueError):
            pass
    for _ in range(4):
        doc.add_paragraph()
    k = doc.add_paragraph().add_run("FTTH SURVEY AND DESIGN REPORT")
    k.bold, k.font.size, k.font.color.rgb = True, Pt(11), GOLD
    t = doc.add_paragraph().add_run(clean(d.info.name or "FTTH project"))
    t.bold, t.font.size, t.font.color.rgb = True, Pt(24), NAVY
    sub = " | ".join(x for x in (d.info.area, d.info.client, d.info.cluster) if x)
    if sub:
        s = doc.add_paragraph().add_run(clean(sub))
        s.font.size, s.font.color.rgb = Pt(12), GREY
    doc.add_paragraph()
    _heading(doc, "Document control", 2)
    _kv(doc, report.document_control(d, st))
    _heading(doc, "Sign-off", 2)
    _table(doc, ["Role", "Name", "Signature", "Date"], [[k, v or "", "", ""] for k, v in report.signoff(d)])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for sec in report.sections(d):
        _heading(doc, sec.title, 1)
        for b in sec.blocks:
            kind = b[0]
            if kind == "p":
                doc.add_paragraph(b[1]).paragraph_format.space_after = Pt(6)
            elif kind == "missing":
                r = doc.add_paragraph().add_run(b[1] + " (to be completed in the project details)")
                r.italic, r.font.color.rgb = True, RGBColor(0x9A, 0x3B, 0x3B)
            elif kind == "h":
                _heading(doc, b[1], 2)
            elif kind == "kv":
                _kv(doc, b[1])
            elif kind == "table":
                _table(doc, b[1], b[2], b[3])
            elif kind == "cards":
                tb = doc.add_table(rows=0, cols=3)
                for i in range(0, len(b[1]), 3):
                    cells = tb.add_row().cells
                    for j, (label, value) in enumerate(b[1][i:i + 3]):
                        cells[j].text = ""
                        v = cells[j].paragraphs[0].add_run(value)
                        v.bold, v.font.size, v.font.color.rgb = True, Pt(15), NAVY
                        lab = cells[j].add_paragraph().add_run(label.upper())
                        lab.font.size, lab.font.color.rgb = Pt(7.5), GREY
                        _shade(cells[j], "EEF3FA")
                doc.add_paragraph()
            elif kind == "bullets":
                for x in b[1]:
                    doc.add_paragraph(x, style="List Bullet")

    photos = [(ph, photo_path(photo_dir, ph.file)) for ph in d.photos]
    photos = [(ph, path) for ph, path in photos if path]
    if photos:
        doc.add_page_break()
        _heading(doc, "Appendix A. Survey Photographs", 1)
        for ph, path in photos:
            try:
                doc.add_picture(str(path), width=Cm(11))
            except (OSError, ValueError):
                continue
            cap = doc.add_paragraph()
            r = cap.add_run(clean(ph.caption or ph.file) + (clean(f" ({ph.feature})") if ph.feature else ""))
            r.bold, r.font.size = True, Pt(9)
            a = ph.analysis or {}
            if a.get("summary") or a.get("objects"):
                ai = cap.add_run(clean("  Photo AI: " + (a.get("summary") or ", ".join(o.get("label", "") for o in a["objects"][:6]))
                                       + (f" Issues: {', '.join(a['issues'])}." if a.get("issues") else "")))
                ai.font.size, ai.font.color.rgb = Pt(9), BLUE
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
