"""Write the demo sample project to samples/: a survey pack to import, the finished design, and its deliverables.

All data is the built-in demo project (ftth.sample: "Demo estate, Nairobi", made up for trying the tool).
Usage: python tools/make_samples.py
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
sys.path.insert(0, str(ROOT))
from blind_test import survey_only  # noqa: E402
from ftth.exporters import excel, kml, package  # noqa: E402
from ftth.exporters.common import Brand, stamp  # noqa: E402
from ftth.sample import sample_design  # noqa: E402

LOGO = ROOT / "planner" / "static" / "planner" / "img" / "archplanner-logo.png"


def survey_workbook(d) -> bytes:
    """The Excel upload template filled with what a survey team delivers: poles, premises, existing network."""
    s = survey_only(d)
    wb = load_workbook(io.BytesIO(excel.survey_template_xlsx()))
    for p in s.poles:
        wb["Poles"].append([p.id, p.lat, p.lon, p.type, p.height_m, p.condition, p.status, p.notes])
    for h in s.homes:
        kind = {"mdu": "MDU", "business": "Business", "home": "SDU"}[h.kind]
        wb["Homes"].append([h.id, h.lat, h.lon, h.name, kind, h.units, h.address, "yes" if h.acquired else "",
                            h.field_type, "", h.access_notes, h.notes])
    for n in s.nodes:
        wb["Equipment"].append([n.id, n.kind, n.lat, n.lon, n.pole, n.status, n.ports, n.notes])
    for c in s.cables:
        wb["Cables"].append([c.id, c.kind, c.cores, "", "", " > ".join(v.pole or "pt" for v in c.path), c.install,
                             c.status, c.notes])
        wb["Cables"].cell(row=wb["Cables"].max_row, column=10, value="; ".join(f"{v.lat:.7f},{v.lon:.7f}" for v in c.path))
    wb["Cables"].cell(row=1, column=10, value="Path coordinates")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main(out: Path = ROOT / "samples") -> None:
    d = sample_design()
    st = stamp(d, 1, Brand(name="ArchPlanner", logo=LOGO if LOGO.is_file() else None))
    files = {
        "1-survey/demo-survey-data.xlsx": survey_workbook(d),
        "2-design/demo-design.kmz": kml.build_kmz(d, st),
        "2-design/demo-project.json": json.dumps(d.model_dump(mode="json"), indent=1).encode(),
        "3-deliverables/demo-design-package.zip": package.build_package(d, st, None),
    }
    for name, data in files.items():
        path = out / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
        print(f"{path.relative_to(ROOT)}  {len(data) / 1024:,.0f} KB")


if __name__ == "__main__":
    main()
