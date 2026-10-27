"""Turn a project folder (KMZ [+ home list]) into a survey-only pack: the Excel upload template filled with
what a survey team delivers (existing poles, existing network, premises with HP counts) and no design.

Usage: python tools/make_survey_pack.py CALIBRATION_FOLDER OUT_FOLDER
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from blind_test import survey_only  # noqa: E402
from calibrate import load_design  # noqa: E402
from ftth.exporters import excel  # noqa: E402


def main(folder: str, out: str):
    d, _ = load_design(Path(folder))
    s = survey_only(d)
    wb = load_workbook(io.BytesIO(excel.survey_template_xlsx()))
    for p in s.poles:
        wb["Poles"].append([p.id, p.lat, p.lon, p.type, p.height_m, p.condition, p.status, p.notes])
    for h in s.homes:
        kind = {"mdu": "MDU", "business": "Business", "home": "SDU"}[h.kind]
        wb["Homes"].append([h.id, h.lat, h.lon, h.name, kind, h.units, h.address, "yes" if h.acquired else "",
                            h.field_type, "", h.access_notes, h.notes])
    for n in s.nodes:
        wb["Equipment"].append([n.id, n.kind, n.lat, n.lon, n.pole, n.status, n.ports, n.notes])
    for c in s.cables:
        wb["Cables"].append([c.id, c.kind, c.cores, "", "", " > ".join(v.pole or "pt" for v in c.path), c.install,
                             c.status, c.notes, ])
        wb["Cables"].cell(row=wb["Cables"].max_row, column=10, value="; ".join(f"{v.lat:.7f},{v.lon:.7f}" for v in c.path))
    wb["Cables"].cell(row=1, column=10, value="Path coordinates")
    o = Path(out)
    o.mkdir(parents=True, exist_ok=True)
    wb.save(o / "survey-data.xlsx")
    brief = {
        "area": d.info.area or d.info.name, "existing_equipment": [n.id for n in s.nodes],
        "backbone_extension_points (lat, lon)": s.info.extend_to,
        "counts": {"poles": len(s.poles), "premises": len(s.homes), "homepasses": sum(h.units for h in s.homes),
                   "existing_cables": len(s.cables)},
    }
    (o / "survey-brief.json").write_text(json.dumps(brief, indent=2))
    print(json.dumps(brief, indent=2))


if __name__ == "__main__":
    main(*sys.argv[1:3])
