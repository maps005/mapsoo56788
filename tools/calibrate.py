"""Compare the planner's quantities with a completed project's BOM.

Usage:  .venv\\Scripts\\python tools\\calibrate.py calibration\\infinity-fy26-approved
The folder must contain design.kmz and BOM.xlsx. Two V Squared BOM layouts are read:
draft ('Pole Details', 'CABLE LENGTH', 'BOM') and approved ('POLE SUMMARY', 'B.O.M').
Poles are matched by position (names can change between revisions); an optional
home-list.xlsx is imported for exact homepass counts.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from ftth import calc, importers  # noqa: E402
from ftth.geo import haversine  # noqa: E402
from ftth.model import DEFAULT_MATERIALS, Design  # noqa: E402

BOM_CODES = {  # BOM wording (draft and approved) -> planner material code
    "12C ADSS FOC": "FOC-12C", "48C ADSS FOC": "FOC-48C", "24C ADSS FOC (20M -TIE CABLES)": "TIE-24C",
    "VSQ 8M POLES": "POLE-NEW", "TENSION CLAMPS": "TENSION-CLAMP", "UPB": "UPB", "J-HOOK": "J-HOOK",
    "CABLE TIES": "CABLE-TIES", "WALL CLAMP": "WALL-CLAMP", "BUCKLES": "BUCKLE-CLIP",
    "PLASTIC TRUNKING SIZE 38*25MM": "TRUNKING", "STAINLESS STEEL STAPS": "STEEL-STRAP",
    "24P INDOOR FAT": "FAT-24P", "16P OUTDOOR FAT": "FAT-16P", "SPLICE CLOSURE": "ENC-SPLICE",
    "JUNCTION BOX": "JB", "1X8 SPLITTER": "SPL-1X8", "1X16 SPLITTER": "SPL-1X16",
    "ADSS FIBER CABLE, 100M S/S CLEAR SPAN 12 XSM G652D": "FOC-12C",
    "ADSS FIBER CABLE, 100M S/S CLEAR SPAN 48 XSM G652D": "FOC-48C",
}


def _num(v):
    if isinstance(v, (int, float)):
        return v
    m = re.search(r"\d+(\.\d+)?", str(v or ""))
    return float(m.group()) if m else None


def load_design(folder: Path) -> tuple[Design, dict]:
    d = Design(materials=[m.model_copy() for m in DEFAULT_MATERIALS])
    report = importers.import_file(d, "design.kmz", (folder / "design.kmz").read_bytes(), "auto")
    hl = folder / "home-list.xlsx"
    if hl.exists():
        r2 = importers.import_file(d, hl.name, hl.read_bytes(), "homes")
        report["problems"] += ["Home list: " + p for p in r2["problems"][:3]]
    return d, report


def reference(folder: Path) -> dict:
    wb = load_workbook(folder / "BOM.xlsx", data_only=True)
    poles, cables, bom = {}, [], {}
    by_desc = {m.description.upper(): m.code for m in DEFAULT_MATERIALS} | BOM_CODES
    if "POLE SUMMARY" in wb.sheetnames:  # approved layout
        ws = wb["POLE SUMMARY"]
        head = [str(c or "").strip().upper() for c in next(ws.iter_rows(values_only=True))]
        col = {h: i for i, h in enumerate(head)}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0] and str(r[0]).strip().upper() != "TOTAL":
                a, b = r[col["LATITUDE"]], r[col["LONGITUDE"]]
                lat, lon = (b, a) if abs(a) > abs(b) else (a, b)  # the sheet's lat/lon headers are swapped
                poles[str(r[0]).strip()] = {"lat": lat, "lon": lon, "tension": r[col["TENSION CLAMPS"]] or 0,
                                            "upb": r[col["UPBS"]] or 0, "jhooks": r[col["J-HOOK"]] or 0,
                                            "ext_arms": 0}
        for r in wb["B.O.M"].iter_rows(values_only=True):
            cells = [c for c in r if c is not None]  # ITEM, UOM, QTY (column positions vary)
            if len(cells) >= 3:
                name, q = re.sub(r"\s+", " ", str(cells[0]).strip().upper()), _num(cells[2])
                if name in by_desc and q is not None:
                    bom[by_desc[name]] = q
    else:  # draft layout
        for r in wb["Pole Details"].iter_rows(min_row=2, values_only=True):
            if r[0] and r[0] != "Total":
                poles[r[0].strip()] = {"lat": r[4], "lon": r[3], "tension": r[6] or 0, "upb": r[7] or 0,
                                       "jhooks": r[8] or 0, "ext_arms": r[9] or 0}
        cables = [r for r in wb["CABLE LENGTH"].iter_rows(min_row=2, values_only=True)
                  if r[0] and isinstance(r[1], (int, float)) and r[2] is not None]
        for r in wb["BOM"].iter_rows(min_row=2, values_only=True):
            if r[2] and r[6] is not None:
                bom[by_desc.get(r[2].strip().upper(), r[2].strip())] = r[6]
    if "TIE-24C" in bom:  # 24C tie cables are counted in 20 m pieces
        bom["FOC-24C"] = bom.pop("TIE-24C") * 20
    return {"poles": poles, "cables": cables, "bom": bom}


def compare(folder: Path, verbose: bool = True) -> dict:
    d, report = load_design(folder)
    ref = reference(folder)
    ours_bom = {l.code: l.qty for l in calc.bom(d)}
    rows = [(code, ref["bom"].get(code), ours_bom.get(code)) for code in sorted(set(ref["bom"]) | set(ours_bom))]

    ours_poles = {p.pole: p for p in calc.pole_summary(d)}
    pos = {p.id: (p.lat, p.lon) for p in d.poles}
    match = {}  # BOM pole name -> planner pole id (nearest within 3 m)
    for name, r in ref["poles"].items():
        best = min(((haversine(r["lat"], r["lon"], *xy), pid) for pid, xy in pos.items() if xy[0] is not None),
                   default=None)
        if best and best[0] <= 3:
            match[name] = best[1]
    keys = ("tension", "upb", "jhooks", "ext_arms")
    pole_diff = []
    for name, r in ref["poles"].items():
        o = ours_poles.get(match.get(name, ""))
        got = {k: getattr(o, k) for k in keys} if o else None
        want = {k: r[k] for k in keys}
        if not r.get("ext_arms") and got:
            got["ext_arms"] = want["ext_arms"] = 0 if "ext_arms" not in r else got["ext_arms"]
        if got != want:
            pole_diff.append((name, match.get(name), want, got, o.cables if o else None))
    matched = set(match.values())
    on_ref = {k: sum(getattr(ours_poles[p], k) for p in matched if p in ours_poles) for k in keys}
    ref_tot = {k: sum(r[k] for r in ref["poles"].values()) for k in keys}
    off_ref = {k: sum(getattr(r, k) for p, r in ours_poles.items() if p not in matched) for k in keys}
    exact = len(ref["poles"]) - len(pole_diff)
    lengths = calc.cable_lengths(d, proposed_only=True)
    if verbose:
        print("\n".join(report["problems"][:5]))
        print(f"\nHomepasses: {calc.summary(d)['homepasses']}")
        print(f"Pole accessories on the {len(ref['poles'])} BOM poles ({len(match)} matched by position): "
              f"planner {on_ref} vs BOM {ref_tot}; {exact} poles identical")
        print(f"Accessories on poles the BOM does not list: {off_ref}\n")
        print(f"{'Item':16s} {'BOM':>8s} {'Planner':>8s}  diff")
        for code, a, b in rows:
            diff = "" if a == b else ("  <-- " + (f"{(b or 0) - (a or 0):+g}"))
            print(f"{code:16s} {a if a is not None else '-':>8} {b if b is not None else '-':>8}{diff}")
        for name, pid, want, got, cabs in pole_diff:
            print(f"  {name:14s} = {pid or '?':12s} BOM {want}  planner {got}  cables {cabs}")
        if ref["cables"]:
            print("\nCable lengths (BOM order vs planner order):")
            for i, cl in enumerate(lengths):
                rc = ref["cables"][i] if i < len(ref["cables"]) else None
                print(f"  {cl.cable:6s} {cl.cores}C measured {cl.route_m:7.1f} / {rc[1] if rc else '-':>7}  "
                      f"total {cl.total_m:5.0f} / {rc[8] if rc else '-'}")
    return {"bom": rows, "pole_diff": pole_diff, "design": d, "on_ref": on_ref, "ref_tot": ref_tot,
            "off_ref": off_ref, "exact": exact, "matched": len(match)}


if __name__ == "__main__":
    compare(Path(sys.argv[1] if len(sys.argv) > 1 else ROOT / "calibration" / "infinity-fy26-approved"))
